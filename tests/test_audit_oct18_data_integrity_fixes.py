# -*- coding: utf-8 -*-
"""
test_audit_oct18_data_integrity_fixes.py - 针对 4 个关键数据完整性问题的回归测试套件：
1. [P1] MinerU 跨格式合并坐标单位统一 (norm_1000 vs points)
2. [P1] 再次导出保护磁盘上已有冲突版本 (_v2.xlsx 避让与 _v3.xlsx 递增)
3. [P1] 记录保留率样本身份与字段对应关系核验 (防止样本丢失)
4. [P1] 冲突检测识别样本间数值互换与数据冲突 (杜绝全表单元格集合掩盖样本值冲突)
"""

import io
import json
import os
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import sys

SCRIPT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "scripts")
if SCRIPT_DIR not in sys.path:
    sys.path.insert(0, SCRIPT_DIR)

import pandas as pd
from openpyxl import load_workbook

import excel_export as ee
import ocr_client as oc
import pdf_table_extractor as pe
import table_postprocess as tp


class TestDataIntegrityFixes(unittest.TestCase):

    def setUp(self):
        self.tmpdir = tempfile.mkdtemp()

    def tearDown(self):
        if os.path.exists(self.tmpdir):
            shutil.rmtree(self.tmpdir, ignore_errors=True)

    # -------------------------------------------------------------------------
    # Issue 1: MinerU 坐标单位统一 (norm_1000 vs points)
    # -------------------------------------------------------------------------
    def test_mineru_coordinate_fusion_unifies_norm1000_and_points(self):
        """[P1] MinerU 跨格式合并统一归一化至 0-1000 空间，绝不误删独立表格"""
        def make_html(rows):
            return '<table><tr><th>Sample</th><th>Value</th></tr>' + ''.join(
                '<tr>' + ''.join(f'<td>{c}</td>' for c in r) + '</tr>' for r in rows
            ) + '</table>'

        def make_archive(files_dict):
            bio = io.BytesIO()
            with zipfile.ZipFile(bio, 'w') as z:
                for fname, d in files_dict.items():
                    z.writestr(fname, json.dumps(d))
            return bio.getvalue()

        t_a = make_html([['Alpha', 13], ['Beta', 25]])
        t_b = make_html([['Gamma', 31], ['Delta', 46], ['Epsilon', 59]])

        # content_list 使用 0-1000 归一化坐标 [100, 500, 900, 625]
        v1 = [{'type': 'table', 'table_body': t_a, 'page_idx': 0, 'bbox': [100, 500, 900, 625]}]

        # middle.json 使用 PDF 点坐标，页面尺寸 [600, 800]
        # 表1 点坐标 [60, 400, 540, 500] -> 归一化为 [100, 500, 900, 625] (与 v1 一致)
        # 表2 点坐标 [100, 500, 550, 625] -> 归一化为 [166.67, 625, 916.67, 781.25] (与表1在垂直方向不重合)
        mid = {
            'pdf_info': [{
                'page_idx': 0,
                'page_size': [600, 800],
                'preproc_blocks': [
                    {'type': 'table', 'bbox': [60, 400, 540, 500], 'blocks': [{'type': 'table_body', 'lines': [{'spans': [{'type': 'table', 'html': t_a}]}]}]},
                    {'type': 'table', 'bbox': [100, 500, 550, 625], 'blocks': [{'type': 'table_body', 'lines': [{'spans': [{'type': 'table', 'html': t_b}]}]}],
                }]
            }]
        }

        # 仅解析 middle.json: 得到 2 张表
        res_mid_only = oc.extract_mineru_structured_archive(make_archive({'middle.json': mid}))
        self.assertEqual(len(res_mid_only), 2, "单独 middle.json 应包含 2 张表")

        # 同时提供 content_list.json 与 middle.json: 必须统一坐标，不能因为混用单位将表2误判为表1的重复而删除
        res_combined = oc.extract_mineru_structured_archive(make_archive({
            'content_list.json': v1,
            'middle.json': mid,
        }))
        self.assertEqual(len(res_combined), 2, "融合 content_list 与 middle.json 时，2 张表必须全部保留")

        all_samples = []
        for df in res_combined:
            if 'Sample' in df.columns:
                all_samples.extend(df['Sample'].tolist())
        self.assertIn('Alpha', all_samples)
        self.assertIn('Beta', all_samples)
        self.assertIn('Gamma', all_samples)
        self.assertIn('Delta', all_samples)
        self.assertIn('Epsilon', all_samples)

    # -------------------------------------------------------------------------
    # Issue 2: 再次导出会覆盖此前保存的冲突版本
    # -------------------------------------------------------------------------
    def test_subsequent_export_does_not_overwrite_existing_v2_version(self):
        """[P1] 多次调用 save_tables_to_excel 时保护磁盘既有 _v2.xlsx，自增至 _v3.xlsx"""
        folder = Path(self.tmpdir) / "version_protection"
        folder.mkdir(parents=True, exist_ok=True)

        t1 = pd.DataFrame({'Sample': ['A', 'B'], 'Value': ['13', '25']})
        t1.attrs.update({'label': 'Table 1', 'table_title': 'Table 1', 'page_idx': 0})

        t2 = pd.DataFrame({'Sample': ['X', 'Y'], 'Value': ['31', '46']})
        t2.attrs.update({'label': 'Table 1', 'table_title': 'Table 1', 'page_idx': 0})

        t3 = pd.DataFrame({'Sample': ['U', 'V'], 'Value': ['57', '69']})
        t3.attrs.update({'label': 'Table 1', 'table_title': 'Table 1', 'page_idx': 0})

        # 第一次导出：产生 Table 1.xlsx 与 Table 1_v2.xlsx
        ee.save_tables_to_excel([t1, t2], str(folder), single_file=False)
        self.assertTrue((folder / 'Table 1.xlsx').exists())
        self.assertTrue((folder / 'Table 1_v2.xlsx').exists())
        old_v2_bytes = (folder / 'Table 1_v2.xlsx').read_bytes()

        # 第二次导出：传入新的冲突版本 t3
        # 必须保护磁盘已有的 Table 1_v2.xlsx，生成 Table 1_v3.xlsx
        ee.save_tables_to_excel([t3], str(folder), single_file=False)

        new_v2_bytes = (folder / 'Table 1_v2.xlsx').read_bytes()
        self.assertEqual(old_v2_bytes, new_v2_bytes, "磁盘上的 Table 1_v2.xlsx 绝不能被第二次导出覆写")
        self.assertTrue((folder / 'Table 1_v3.xlsx').exists(), "新的冲突版本必须自增另存为 Table 1_v3.xlsx")

        wb_v3 = load_workbook(folder / 'Table 1_v3.xlsx')
        v3_values = list(wb_v3.active.values)
        wb_v3.close()
        self.assertEqual(v3_values[1], ('U', 57))
        self.assertEqual(v3_values[2], ('V', 69))

    # -------------------------------------------------------------------------
    # Issue 3: “记录保留率”核查样本身份与字段对应关系
    # -------------------------------------------------------------------------
    def test_record_retention_checks_sample_identity_and_field_correspondence(self):
        """[P1] 记录保留率必须核实样本身份与字段对应，杜绝 OCR 丢弃样本依然准予替换"""
        clean = pd.DataFrame({
            'Sample': [f'S{i}' for i in range(10)],
            'Si': ['13'] * 10,
            'Ca': ['37'] * 10
        })
        clean.attrs.update({'label': 'Table 1', 'table_title': 'Table 1', 'page_idx': 0})

        # 原生表包含 10 个样本 + 2 行元数据污染行
        nat = pd.concat([
            clean,
            pd.DataFrame({'Sample': ['received: 2020', 'accepted: 2021'], 'Si': ['', ''], 'Ca': ['', '']})
        ], ignore_index=True)
        nat.attrs.update(clean.attrs)

        # 反例 1：OCR 仅提取 S0-S4，重复 2 遍凑成 10 行，丢失 S5-S9
        ocr_repeated_half = pd.concat([clean.iloc[:5], clean.iloc[:5]], ignore_index=True)
        ocr_repeated_half.attrs.update(clean.attrs)

        should_replace = pe._should_replace_with_ocr({'df': nat, 'page_idx': 0}, {'df': ocr_repeated_half, 'page_idx': 0})
        self.assertFalse(should_replace, "OCR 仅包含 S0~S4 重复、丢失 S5~S9 时绝不能通过替换门禁")

        # 反例 2：OCR 提取 S0-S8，末行重复 S0，丢失 S9
        ocr_drop_one = pd.concat([clean.iloc[:9], clean.iloc[:1]], ignore_index=True)
        ocr_drop_one.attrs.update(clean.attrs)

        should_replace_2 = pe._should_replace_with_ocr({'df': nat, 'page_idx': 0}, {'df': ocr_drop_one, 'page_idx': 0})
        self.assertFalse(should_replace_2, "OCR 丢失 S9 时绝不能通过替换门禁")

        # 正例：OCR 完整保留 S0~S9 全部 10 个样本，剔除污染行
        ocr_full = clean.copy()
        ocr_full.attrs.update(clean.attrs)

        should_replace_full = pe._should_replace_with_ocr({'df': nat, 'page_idx': 0}, {'df': ocr_full, 'page_idx': 0})
        self.assertTrue(should_replace_full, "OCR 真实完整保留所有 S0~S9 样本时应获准替换")

    def test_record_retention_protects_uncontaminated_native_table(self):
        """[P1] 即使原生表未被标记为污染或低质，OCR 包含重复假样本并丢失 S5~S9 时也绝不能替换原生表"""
        clean = pd.DataFrame({
            'Sample': [f'S{i}' for i in range(10)],
            'Si': ['13'] * 10,
            'Ca': ['37'] * 10
        })
        clean.attrs.update({'label': 'Table 1', 'table_title': 'Table 1', 'page_idx': 0})

        ocr_dup = pd.concat([clean.iloc[:5], clean.iloc[:5]], ignore_index=True)
        ocr_dup.attrs.update({'label': 'Table 1', 'table_title': 'Table 1', 'page_idx': 0})

        # 验证门禁拒绝
        self.assertFalse(pe._should_replace_with_ocr({'df': clean}, {'df': ocr_dup}),
                         "清洁原生表绝不能被丢失 S5~S9 的重复 OCR 替换")

        # 验证融合调和引擎保留原生 10 行并记录日志
        logs = []
        reconciled = pe.reconcile_native_and_ocr_tables(
            [{'df': clean, 'page_idx': 0, 'label': 'Table 1'}],
            [{'df': ocr_dup, 'page_idx': 0, 'label': 'Table 1'}],
            [0],
            logs
        )
        self.assertEqual(len(reconciled), 1)
        res_df = reconciled[0]['df']
        self.assertEqual(len(res_df), 10, "原生 10 个样本必须完整保留")
        self.assertEqual(res_df['Sample'].tolist(), [f'S{i}' for i in range(10)])
        self.assertTrue(any('保留原生版本' in l for l in logs))

    # -------------------------------------------------------------------------
    # Issue 4: 冲突检测识别数值在样本间互换与数据冲突
    # -------------------------------------------------------------------------
    def test_conflict_detection_identifies_swapped_values_across_samples(self):
        """[P1] 冲突检测必须比对样本—字段—值关联，全表单元格相同但样本间数值互换判定为冲突"""
        def make_table(data, page=0, source=None):
            d = pd.DataFrame(data)
            d.attrs.update({'label': 'Table 1', 'table_title': 'Table 1', 'page_idx': page})
            if source:
                d.attrs['source'] = source
            return d

        # 表1: A=13, B=25
        # 表2: A=25, B=13 (数值在样本间互换，单元格集合完全相同)
        df_a = make_table({'Sample': ['A', 'B'], 'Si': ['13', '25']}, source='online')
        df_b = make_table({'Sample': ['A', 'B'], 'Si': ['25', '13']}, source='pdf')

        is_conflict = tp.is_conflicting_or_distinct_version(df_a, df_b)
        self.assertTrue(is_conflict, "样本间数值互换属于严重数据冲突，必须判定为 conflicting 独立版本")

        # 验证单文件 Excel 导出时双方均保留为独立 Sheet，不发生覆盖丢弃
        out_file = os.path.join(self.tmpdir, "swapped_export.xlsx")
        ee.save_tables_to_excel([df_a, df_b], out_file, single_file=True)

        wb = load_workbook(out_file)
        self.assertEqual(len(wb.sheetnames), 2, "冲突表格应保存为两个独立工作表")
        self.assertIn("Table 1", wb.sheetnames)
        self.assertTrue(any("pdf" in s.lower() or "v2" in s.lower() for s in wb.sheetnames))

        ws1_data = {row[0]: row[1] for row in list(wb[wb.sheetnames[0]].values)[1:]}
        ws2_data = {row[0]: row[1] for row in list(wb[wb.sheetnames[1]].values)[1:]}
        wb.close()

        self.assertEqual(ws1_data['A'], 13)
        self.assertEqual(ws1_data['B'], 25)
        self.assertEqual(ws2_data['A'], 25)
        self.assertEqual(ws2_data['B'], 13)

    def test_conflict_detection_identifies_swapped_values_in_large_table(self):
        """[P1] 即使大表格（15 样本）冲突率较低，样本间数值互换必须 100% 检出冲突"""
        df1 = pd.DataFrame({'Sample': [f'S{i}' for i in range(15)], 'Val': [str(i * 10) for i in range(15)]})
        df2 = df1.copy()
        # 互换 S0 与 S1
        df2.loc[0, 'Val'] = '10'
        df2.loc[1, 'Val'] = '0'

        self.assertTrue(tp.is_conflicting_or_distinct_version(df1, df2), "15 样本中的数值互换必须检出冲突")

    def test_conflict_detection_unlabelled_columns_and_different_width(self):
        """[P1] 无显式列名的 Col_1, Col_2 表格在列数不同时仍通过位置匹配检出数值互换冲突"""
        df1 = pd.DataFrame({'Col_1': ['A', 'B'], 'Col_2': ['13', '25']})
        df2 = pd.DataFrame({'Col_1': ['A', 'B'], 'Col_2': ['25', '13'], 'Col_3': ['x', 'y']})

        self.assertTrue(tp.is_conflicting_or_distinct_version(df1, df2), "Col_ 命名的不同宽度表格数值互换必须检出冲突")

    def test_mineru_middle_json_preserves_bbox_and_caption_metadata(self):
        """[P1] MinerU middle.json 结构树提取时完整保留表格 bbox、caption 与 footnote 元数据"""
        def make_archive(files_dict):
            bio = io.BytesIO()
            with zipfile.ZipFile(bio, 'w') as z:
                for fname, d in files_dict.items():
                    z.writestr(fname, json.dumps(d))
            return bio.getvalue()

        table_html = '<table><tr><th>Sample</th><th>Value</th></tr><tr><td>Alpha</td><td>13</td></tr></table>'
        mid = {
            'pdf_info': [{
                'page_idx': 0,
                'page_size': [600, 800],
                'preproc_blocks': [{
                    'type': 'table',
                    'bbox': [60, 400, 540, 500],
                    'blocks': [
                        {'type': 'table_caption', 'lines': [{'spans': [{'type': 'text', 'content': 'Table 1. Experimental Data'}]}]},
                        {'type': 'table_body', 'lines': [{'spans': [{'type': 'table', 'html': table_html}]}]},
                        {'type': 'table_footnote', 'lines': [{'spans': [{'type': 'text', 'content': '*P < 0.05'}]}]}
                    ]
                }]
            }]
        }
        dfs = oc.extract_mineru_structured_archive(make_archive({'middle.json': mid}))
        self.assertEqual(len(dfs), 1)
        df = dfs[0]
        self.assertEqual(df.attrs.get('bbox'), [60, 400, 540, 500], "bbox 必须从 middle.json table 块完整保留")
        self.assertEqual(df.attrs.get('table_title'), 'Table 1. Experimental Data')
        self.assertEqual(df.attrs.get('label'), 'Table 1')
        self.assertEqual(df.attrs.get('table_footnote'), '*P < 0.05')

    def test_conflict_detection_negative_case_identical_and_extended_tables(self):
        """[P1] 负例测试：真正完全相同或单纯扩展列的兼容表格不应误报冲突"""
        t1 = pd.DataFrame({'Sample': ['A', 'B'], 'Si': ['13', '25']})
        t2 = pd.DataFrame({'Sample': ['A', 'B'], 'Si': ['13', '25']})
        self.assertFalse(tp.is_conflicting_or_distinct_version(t1, t2), "相同内容表格不构成冲突")

        t3 = pd.DataFrame({'Sample': ['A', 'B'], 'Si': ['13', '25'], 'Ca': ['7', '8']})
        self.assertFalse(tp.is_conflicting_or_distinct_version(t1, t3), "相同样本具有一致测量值的扩展列表格不构成冲突")


if __name__ == '__main__':
    unittest.main()
