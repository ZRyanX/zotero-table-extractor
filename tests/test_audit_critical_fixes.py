import os
import sys
import unittest
import tempfile
import pathlib
import types
from unittest.mock import patch, MagicMock
import pandas as pd
import pymupdf as fitz
from PIL import Image

sys.path[:0] = [
    os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "scripts")),
    os.path.abspath(os.path.join(os.path.dirname(__file__), "..")),
]

import table_postprocess as tp
import excel_export as ex
import pdf_table_extractor as pte
import pdf_tables
import common
import doc_orientation_detector as ori
from omnidocbench import compute_table_teds, compute_layout_detection_map, compute_omission_recall
from rapid_table_extractor import RapidTableExtractor
from teds import TEDS, teds_similarity
from online.strategies import parse_tables_from_html


class TestAuditCriticalFixes(unittest.TestCase):
    def test_integer_rows_not_deleted(self):
        """P1-1: 包含纯整数的合法数据行不应被当成列索引行删除"""
        integers = pd.DataFrame([[10, 20, 30], [40, 50, 60]], columns=['A', 'B', 'C'])
        processed = tp.postprocess_dataframe(integers)
        self.assertEqual(processed.shape, (2, 3))
        self.assertEqual(processed.iloc[0].tolist(), [10, 20, 30])
        self.assertEqual(processed.iloc[1].tolist(), [40, 50, 60])

    def test_single_row_anonymous_columns_protected(self):
        """P1-1: 单行表且使用默认列名 (Col_0/Col_1) 时，有效数据列不应被全部删除"""
        anonymous = pd.DataFrame([[1.5, 2.5]], columns=['Col_0', 'Col_1'])
        processed = tp.postprocess_dataframe(anonymous)
        self.assertEqual(processed.shape, (1, 2))
        self.assertEqual(processed.iloc[0].tolist(), [1.5, 2.5])

    def test_explicit_index_row_filtered_in_continuation(self):
        """P1-1: 带有括号的列序号行 (1) (2) (3) 应被正确识别为索引行"""
        base = pd.DataFrame([['A', 'B', 'C']], columns=['Col_1', 'Col_2', 'Col_3'])
        base.attrs = {"label": "Table 1", "page_idx": 0}
        next_part = pd.DataFrame([['(1)', '(2)', '(3)'], ['D', 'E', 'F']], columns=['Col_1', 'Col_2', 'Col_3'])
        next_part.attrs = {"label": "Table 1", "page_idx": 1}
        merged = tp.merge_continuation_tables([base, next_part])
        self.assertEqual(len(merged), 1)
        self.assertEqual(len(merged[0]), 2)
        self.assertEqual(merged[0].iloc[1].tolist(), ['D', 'E', 'F'])

    def test_multiple_crops_per_page_all_recognized(self):
        """P1-2: 同页多张切图应循环逐一识别，不能只识别最后一张"""
        with tempfile.TemporaryDirectory() as td:
            pdf_path = os.path.join(td, "test_crops.pdf")
            doc = fitz.open()
            doc.new_page()
            doc.save(pdf_path)
            doc.close()

            crops = [
                {'page_index': 0, 'image': Image.new('RGB', (64, 64), 'red'), 'crop_bbox': [0, 0, 64, 64], 'pdf_bbox': [0, 0, 64, 64]},
                {'page_index': 0, 'image': Image.new('RGB', (64, 64), 'blue'), 'crop_bbox': [100, 100, 164, 164], 'pdf_bbox': [100, 100, 164, 164]}
            ]
            called_colors = []

            def mock_api(path, **kwargs):
                with Image.open(path) as im:
                    called_colors.append(im.getpixel((0, 0)))
                return '| Col1 | Col2 |\n|---|---|\n| 1 | 2 |'

            with patch.object(pdf_tables, 'extract_table_crops_from_pdf', return_value=crops), \
                 patch.object(pte.ocr_client, 'call_paddleocr_vl_online_api', side_effect=mock_api), \
                 patch.object(ori, 'detect_crop_orientation', return_value={'needs_rotation': False}), \
                 patch.object(pte, 'verify_and_rescue_table_with_agent', side_effect=lambda df, *args: df), \
                 patch.object(pte, 'extract_via_paddleocr_fullpage', return_value=[]):
                results = pte.extract_via_paddleocr_crops(pdf_path, pages=[0])

            self.assertEqual(len(called_colors), 2)
            self.assertEqual(called_colors, [(255, 0, 0), (0, 0, 255)])
            self.assertEqual(len(results), 2)

    def test_workbook_supplement_preserves_existing_sheets(self):
        """P1-3: 单工作簿模式补漏写入时，必须保留已有的 Table 1，并原子写入 Table 2"""
        with tempfile.TemporaryDirectory() as td:
            out_file = os.path.join(td, "tables.xlsx")
            t1 = pd.DataFrame([['Alpha', 1.5], ['Beta', 2.5]], columns=['Sample', 'Value'])
            t1.attrs.update(label='Table 1', table_title='Table 1')
            ex.save_tables_to_excel([t1], out_file, single_file=True)

            sheets_before = pd.read_excel(out_file, sheet_name=None)
            self.assertIn('Table 1', sheets_before)

            t2 = pd.DataFrame([['Gamma', 3.5], ['Delta', 4.5]], columns=['Sample', 'Value'])
            t2.attrs.update(label='Table 2', table_title='Table 2')

            ex.save_tables_to_excel([t1, t2], out_file, single_file=True, skip_labels={'Table 1'})
            sheets_after = pd.read_excel(out_file, sheet_name=None)

            self.assertIn('Table 1', sheets_after)
            self.assertIn('Table 2', sheets_after)
            self.assertEqual(sheets_after['Table 1'].values.tolist(), [['Alpha', 1.5], ['Beta', 2.5]])
            self.assertEqual(sheets_after['Table 2'].values.tolist(), [['Gamma', 3.5], ['Delta', 4.5]])

    def test_candidate_page_detects_borderless_table_despite_header_line(self):
        """P1-4: 页面有页眉横线但无表格网格时，不能将无框表排除在候选页之外"""
        with tempfile.TemporaryDirectory() as td:
            pdf_path = os.path.join(td, "header_line.pdf")
            doc = fitz.open()
            page = doc.new_page()
            # 绘制一根页眉横线
            shape = page.new_shape()
            shape.draw_line(fitz.Point(50, 50), fitz.Point(550, 50))
            shape.finish(color=(0, 0, 0), width=1)
            shape.commit()
            # 写入模拟文本表格
            page.insert_text(fitz.Point(100, 150), "ColA\tColB\tColC\n10\t20\t30\n40\t50\t60")
            doc.save(pdf_path)
            doc.close()

            # 模拟 find_tables() 找不到网格线，但 find_tables(strategy="text") 能找到
            mock_table = MagicMock()
            mock_table.row_count = 3
            mock_table.col_count = 3
            mock_tabs_text = MagicMock()
            mock_tabs_text.tables = [mock_table]

            orig_find = fitz.Page.find_tables
            def custom_find_tables(*args, **kwargs):
                if kwargs.get('strategy') == 'text':
                    return mock_tabs_text
                return MagicMock(tables=[])

            with patch.object(fitz.Page, 'find_tables', side_effect=custom_find_tables):
                cand_pages, _, _ = pte.get_table_candidate_pages(pdf_path)
                self.assertIn(0, cand_pages)

    def test_direction_model_no_duplicate_softmax(self):
        """P1-5: 文档方向模型输出概率时，不得进行二次 softmax 阻断旋转"""
        det = ori.DocOrientationDetector.__new__(ori.DocOrientationDetector)
        det.output_name = "output"
        det.input_name = "input"
        det._available = True

        # 模拟模型输出 Top-1 概率 0.95 (已 Softmax 归一化)
        prob_output = [0.02, 0.95, 0.02, 0.01]  # 90° class
        mock_session = MagicMock()
        mock_session.run.return_value = [[prob_output]]
        det.session = mock_session

        dummy_img = Image.new('RGB', (100, 100), 'white')
        res = det.predict(dummy_img, min_confidence=0.50)

        self.assertEqual(res['detected_angle'], 90)
        self.assertTrue(res['needs_rotation'])
        self.assertAlmostEqual(res['confidence'], 0.95, places=3)

    def test_rejected_ocr_table_not_appended_back(self):
        """P1-6: 被质量门禁拒绝的 OCR 表，不得在最终兜底中再次被追加回来"""
        # 原生表质量优秀
        df_nat = pd.DataFrame([['A', 1], ['B', 2], ['C', 3], ['D', 4]], columns=['Col1', 'Col2'])
        df_nat.attrs.update(label='Table 1', table_title='Table 1', extractor='native_find_tables')
        nat_item = {'df': df_nat, 'page_idx': 0, 'bbox': [10, 10, 100, 100], 'pdf_bbox': [10, 10, 100, 100]}

        # OCR 表质量低劣（只有1行）
        df_ocr = pd.DataFrame([['A', 1]], columns=['Col1', 'Col2'])
        df_ocr.attrs.update(label='Table 1', table_title='Table 1', extractor='paddleocr_vl_crop')
        ocr_item = {'df': df_ocr, 'page_idx': 0, 'bbox': [10, 10, 100, 100], 'pdf_bbox': [10, 10, 100, 100]}

        # _should_replace_with_ocr 返回 False (原生表胜出)
        with patch.object(pte, '_should_replace_with_ocr', return_value=False):
            merged = pte.merge_extracted_tables([nat_item], [ocr_item], candidate_pages=[0])

        self.assertEqual(len(merged), 1)
        self.assertEqual(merged[0]['df'].attrs.get('extractor'), 'native_find_tables')

    def test_omnidocbench_teds_and_map_accuracy(self):
        """P1-7: TEDS 必须进行真实计算而不是虚假 1.0；mAP 不得混淆为 precision；零声明召回率为 0.0"""
        # 1. 结构完全不同的两张表 TEDS 绝不能是 1.0
        t1_html = "<table><tr><td>A</td></tr></table>"
        t2_html = "<table><tr><td>B</td><td>999</td><td>Extra</td></tr><tr><td>1</td><td>2</td><td>3</td></tr></table>"
        teds_score = compute_table_teds(t1_html, t2_html)
        self.assertLess(teds_score, 0.7)

        # 2. 10 个 GT 只检出 1 个，mAP 绝不能是 1.0
        gt_boxes = [[i * 20, 0, i * 20 + 10, 10] for i in range(10)]
        pred_boxes = [[0, 0, 10, 10]]
        map_metrics = compute_layout_detection_map(pred_boxes, gt_boxes)
        self.assertAlmostEqual(map_metrics['recall@50'], 0.1, places=3)
        self.assertAlmostEqual(map_metrics['mAP@50'], 0.1, places=3)
        self.assertAlmostEqual(map_metrics['mAP@50:95'], 0.1, places=3)

        # 3. 声明数为 0 时，召回率必须是 0.0
        recall_res = compute_omission_recall(0, 0)
        self.assertEqual(recall_res['recall'], 0.0)

        # 4. OmniDocBenchEvaluator 记录中 map_50 必须采用 mAP@50 而非 precision@50
        from omnidocbench import OmniDocBenchEvaluator
        evaluator = OmniDocBenchEvaluator()
        evaluator.add_paper_result("p_test", "academic", 10, 1, pred_boxes=pred_boxes, gt_boxes=gt_boxes)
        self.assertEqual(len(evaluator.records), 1)
        self.assertAlmostEqual(evaluator.records[0]['map_50'], 0.1, places=3)
        self.assertAlmostEqual(evaluator.records[0]['map_50_95'], 0.1, places=3)

    def test_entrypoint_typing_compatibility(self):
        """入口兼容性：Tuple、Set 必须正确导入，不抛 NameError"""
        import extract_zotero_table as ezt
        with tempfile.TemporaryDirectory() as td:
            pdf_path = os.path.join(td, "test_typing.pdf")
            doc = fitz.open()
            doc.new_page()
            doc.save(pdf_path)
            doc.close()
            labels, count = ezt.scan_pdf_table_expected_labels(pdf_path)
            self.assertIsInstance(labels, set)
            self.assertIsInstance(count, int)

    def test_online_single_column_table_supported(self):
        """单列表支持：通用 HTML 解析不得过滤单列表格"""
        html_content = "<table><tr><th>Items</th></tr><tr><td>Alpha</td></tr><tr><td>Beta</td></tr></table>"
        dfs = parse_tables_from_html(html_content)
        self.assertEqual(len(dfs), 1)
        self.assertEqual(dfs[0].shape, (2, 1))

    def test_img2table_coordinates_scaling(self):
        """img2table 坐标应由 200 DPI 像素统一缩放为 72 DPI PDF 点坐标"""
        mock_bbox = MagicMock(x1=200, y1=200, x2=400, y2=400)
        mock_tbl = MagicMock()
        mock_tbl.df = pd.DataFrame([[1, 2]], columns=['A', 'B'])
        mock_tbl.bbox = mock_bbox
        mock_tbl.title = "Table 1"

        mock_doc = MagicMock()
        mock_doc.extract_tables.return_value = {0: [mock_tbl]}

        with tempfile.TemporaryDirectory() as td:
            pdf_path = os.path.join(td, "dummy.pdf")
            doc = fitz.open()
            doc.new_page()
            doc.save(pdf_path)
            doc.close()

            with patch.object(pte, 'is_img2table_available', return_value=True), \
                 patch.dict(sys.modules, {'img2table.document': MagicMock(PDF=MagicMock(return_value=mock_doc)),
                                          'img2table.ocr': MagicMock()}):
                with patch('pdf_table_extractor.fitz.open', return_value=[MagicMock(rect=fitz.Rect(0, 0, 612, 792))]):
                    res = pte.extract_via_img2table(pdf_path, candidate_pages=[0])
                    # 200 * 72 / 200 = 72
                    self.assertEqual(res[0]['pdf_bbox'], [72.0, 72.0, 144.0, 144.0])
                    self.assertEqual(res[0]['bbox'], [200, 200, 400, 400])

    def test_rapid_table_extractor_compatibility(self):
        """RapidTable 官方构造器与 pred_html / pred_htmls 输出对象兼容"""
        class MockOfficialRapidTable:
            def __init__(self, cfg=None):
                pass
            def __call__(self, img):
                class MockResult:
                    pred_html = "<table><tr><th>Header</th></tr><tr><td>Data</td></tr></table>"
                return MockResult()

        with patch.dict(sys.modules, {'rapid_table': types.SimpleNamespace(RapidTable=MockOfficialRapidTable)}):
            extractor = RapidTableExtractor()
            extractor._init_engine()
            self.assertIsNotNone(extractor._table_engine)
            df = extractor.extract_table_from_image(Image.new('RGB', (100, 100), 'white'))
            self.assertIsNotNone(df)
            self.assertEqual(df.shape, (1, 1))


    def test_middle_row_123_and_numeric_matrix_row0_not_deleted(self):
        """P1-1: 位于数据中间的 [1, 2, 3] 及纯数值矩阵首行的 [1, 2, 3] 绝不应被误判删除"""
        # 1. 位于表中间的 [1, 2, 3]
        middle_123 = pd.DataFrame([['A', 'B', 'C'], ['D', 'E', 'F'], [1, 2, 3], ['G', 'H', 'I']], columns=['C1', 'C2', 'C3'])
        proc1 = tp.postprocess_dataframe(middle_123)
        self.assertEqual(proc1.shape, (4, 3))
        self.assertEqual(proc1.iloc[2].tolist(), [1, 2, 3])

        # 2. 纯数值矩阵第一行的 [1, 2, 3]
        num_matrix = pd.DataFrame([[1, 2, 3], [4, 5, 6], [7, 8, 9], [10, 11, 12]], columns=['ColA', 'ColB', 'ColC'])
        proc2 = tp.postprocess_dataframe(num_matrix)
        self.assertEqual(proc2.shape, (4, 3))
        self.assertEqual(proc2.iloc[0].tolist(), [1, 2, 3])

    def test_color_column_not_treated_as_placeholder(self):
        """P1-1: 真实英文列名 'Color' / 'Column' 绝不应被误当成占位列被丢弃或合并"""
        df = pd.DataFrame([['Red', 10], ['Blue', 20]], columns=['Color', 'Value'])
        proc = tp.postprocess_dataframe(df)
        self.assertIn('Color', proc.columns)
        self.assertEqual(proc.shape, (2, 2))

    def test_teds_preserves_raw_html_spans(self):
        """P1-7 / OpenSource: TEDS 评估 DataFrame 时必须优先保留 attrs['raw_html'] 中的 colspan/rowspan"""
        html_with_span = "<table><tr><th colspan=\"2\">Header Span</th></tr><tr><td>A</td><td>B</td></tr></table>"
        df = pd.DataFrame([['A', 'B']], columns=['Header 1', 'Header 2'])
        df.attrs["raw_html"] = html_with_span

        # 与自身 HTML 评估应得到 1.0 (因完整采纳 raw_html 的 colspan 结构)
        score = teds_similarity(df, html_with_span)
        self.assertEqual(score, 1.0)

    def test_layout_map_empty_gt_all_keys_present(self):
        """P1-7: 空 GT 时 compute_layout_detection_map 必须提供完整的度量字典且含 mAP@50:95 且不虚标为满分 (0.0)"""
        res = compute_layout_detection_map([], [])
        self.assertIn("mAP@50:95", res)
        self.assertIn("precision@50", res)
        self.assertIn("recall@50", res)
        self.assertEqual(res["mAP@50:95"], 0.0)

    def test_export_crops_to_excel_preserves_existing_dfs(self):
        """P1-3: export_crops_to_excel 传入 existing_dfs 时必须合并写入，不得覆盖丢失"""
        with tempfile.TemporaryDirectory() as td:
            pdf_path = os.path.join(td, "test.pdf")
            doc = fitz.open()
            doc.new_page()
            doc.save(pdf_path)
            doc.close()

            out_xlsx = os.path.join(td, "crops_merged.xlsx")
            online_df = pd.DataFrame([['Online', 99]], columns=['Source', 'Val'])
            online_df.attrs.update(label='Table 1', table_title='Table 1')

            # 当没有检测到 crops 时，existing_dfs 依然应被安全写入
            ok = pdf_tables.export_crops_to_excel(pdf_path, out_xlsx, crops=[], existing_dfs=[online_df])
            self.assertTrue(ok)
            read_back = pd.read_excel(out_xlsx)
            self.assertEqual(read_back.iloc[0].tolist(), ['Online', 99])

    def test_default_numeric_columns_row_123_not_deleted(self):
        """P1-3: 默认数字列名下，4行以上表格中的真实数据行 [1, 2, 3] 绝不被误判为重复表头删除"""
        df = pd.DataFrame([[10, 20, 30], [1, 2, 3], [40, 50, 60], [70, 80, 90]])
        proc = tp.postprocess_dataframe(df)
        self.assertEqual(len(proc), 4)
        self.assertEqual(proc.iloc[1].tolist(), [1, 2, 3])

    def test_doclayout_official_paddleocr_output_format(self):
        """P1-4: 适配层输入 PaddleOCR 官方字典格式或带 coordinate 字段的检测结果时，正确生成表格区域"""
        from doclayout_yolo_detector import assemble_extended_table_regions
        img = Image.new('RGB', (640, 640), 'white')
        
        # 1. 带 cls_id: 21 与 coordinate
        official_dets = [{'cls_id': 21, 'score': 0.98, 'coordinate': [100.0, 100.0, 400.0, 400.0]}]
        res1 = assemble_extended_table_regions(official_dets, img)
        self.assertEqual(len(res1), 1)

        # 2. 官方 res.boxes 顶层字典
        official_dict = {'res': {'boxes': [{'cls_id': 21, 'label': 'table', 'score': 0.98, 'coordinate': [100.0, 100.0, 400.0, 400.0]}]}}
        res2 = assemble_extended_table_regions(official_dict, img)
        self.assertEqual(len(res2), 1)

    def test_rapid_table_pred_htmls_list_and_raw_html_attrs(self):
        """P1: RapidTable 官方 pred_htmls: List[str] 必须正确解包并保留 df.attrs['raw_html']"""
        html = "<table><tr><th>ColA</th><th>ColB</th></tr><tr><td>1</td><td>2</td></tr></table>"
        extractor = RapidTableExtractor()
        extractor._initialized = True

        # 1. 官方 pred_htmls 列表对象
        mock_output = types.SimpleNamespace(pred_htmls=[html])
        extractor._table_engine = extractor._wired_engine = extractor._wireless_engine = lambda img: mock_output
        df = extractor.extract_table_from_image(Image.new('RGB', (64, 64), 'white'))
        self.assertIsNotNone(df)
        self.assertEqual(df.shape, (1, 2))
        self.assertEqual(df.attrs.get("raw_html"), html)

        # 2. 单行全 <td> 表格 (无 <th>)
        single_row_html = "<table><tr><td>1.5</td><td>2.5</td></tr></table>"
        df_single = extractor.parse_html_to_dataframe(single_row_html)
        self.assertIsNotNone(df_single)
        self.assertEqual(df_single.shape, (1, 2))
        self.assertEqual(df_single.attrs.get("raw_html"), single_row_html)

        # 3. 空 pred_htmls 列表必须安全返回 None
        extractor._table_engine = extractor._wired_engine = extractor._wireless_engine = lambda img: types.SimpleNamespace(pred_htmls=[])
        self.assertIsNone(extractor.extract_table_from_image(Image.new('RGB', (64, 64), 'white')))

    def test_pp_doclayout_chart_not_extracted_as_table(self):
        """P1: PP-DocLayoutV3 类别 3 (chart) 映射为 figure，不应被当作 table 提取"""
        import doclayout_yolo_detector as dl
        import numpy as np

        class MockPPSession:
            def get_inputs(self):
                return [types.SimpleNamespace(name='image', shape=[1, 3, 800, 800]),
                        types.SimpleNamespace(name='scale_factor', shape=[1, 2])]
            def run(self, _, feed):
                # 返回类别 3 (chart)
                return [np.array([[3, 0.99, 50, 50, 100, 100]], dtype=np.float32)]

        detector = dl.PPDocLayoutV3Detector.__new__(dl.PPDocLayoutV3Detector)
        detector._available = True
        detector._session = MockPPSession()

        # 类别 3 必须不产生表格切图
        regions = detector.extract_extended_table_regions(Image.new('RGB', (1280, 800), 'white'))
        self.assertEqual(len(regions), 0)

    def test_pp_only_factory_offline_no_yolo_check(self):
        """P2: 仅选择 PP 引擎时，不得调用 ensure_model_file(model_dir) 或要求 YOLO 存在"""
        import doclayout_yolo_detector as dl
        from pathlib import Path
        with tempfile.TemporaryDirectory() as td:
            pp_file = Path(td) / "pp_doclayoutv3.onnx"
            pp_file.write_bytes(b"0" * (1024 * 1024 + 10))

            with patch.object(dl, 'ensure_model_file', side_effect=FileNotFoundError("YOLO absent and offline")) as mock_ensure, \
                 patch.object(dl, 'PPDocLayoutV3Detector') as mock_pp:
                dl.get_layout_detector('pp_doclayoutv3', model_dir=td)
                self.assertEqual(mock_ensure.call_count, 0)
                self.assertEqual(mock_pp.call_count, 1)

    def test_excel_export_preserves_academic_cell_rich_text(self):
        """P2: Excel 补写与格式自适应必须保留 openpyxl CellRichText 学术上下标"""
        from openpyxl import Workbook, load_workbook
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        from pathlib import Path

        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / 'test_rich.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = 'Table 1'
            ws.append(['Species', 'Value'])
            ws.append([CellRichText(TextBlock(InlineFont(), 'H'), TextBlock(InlineFont(vertAlign='subscript'), '2'), TextBlock(InlineFont(), 'O')), 3])
            wb.save(p)
            wb.close()

            # 补写 Table 2
            d = pd.DataFrame({'Species': ['NaCl'], 'Value': [5]})
            d.attrs['label'] = 'Table 2'
            ex.save_tables_to_excel([d], str(p), skip_labels={'Table 1'})

            # 验证 Table 1 A2 仍为 CellRichText 且包含 subscript
            w = load_workbook(p, rich_text=True)
            v = w['Table 1']['A2'].value
            self.assertIsInstance(v, CellRichText)
            runs = [(x.text, getattr(x.font, 'vertAlign', None)) for x in v]
            self.assertEqual(runs, [('H', None), ('2', 'subscript'), ('O', None)])
            w.close()

    def test_corrupt_or_empty_workbook_triggers_recheck(self):
        """P2: 损坏的 xlsx 或全空 sheet 不得计入 actual_labels，必须触发重检生成任务"""
        import batch_run
        import extract_zotero_table as ezt
        from pathlib import Path

        with tempfile.TemporaryDirectory() as td:
            folder = Path(td) / 'paper'
            folder.mkdir()
            pdf = Path(td) / 'paper.pdf'
            pdf.write_bytes(b'dummy pdf')

            # 1. 损坏的工作簿
            (folder / 'Table 1.xlsx').write_bytes(b'corrupted file')
            records = [{'filename': 'paper.pdf', 'pdf_path': str(pdf), 'parent_title': 'paper'}]

            with patch.object(ezt, 'scan_pdf_table_expected_labels', return_value=({'Table 1'}, 1)):
                tasks = batch_run.generate_batch_tasks(td, records, recheck=True)
                self.assertEqual(len(tasks), 1)

    def test_omnidocbench_teds_preserves_spans_and_handles_empty(self):
        """P2: compute_table_teds 必须保留 raw_html spans，对空输入安全返回 0.0"""
        a = pd.DataFrame([['H1', 'H2'], ['A', 'B']])
        b = a.copy()
        a.attrs['raw_html'] = '<table><tr><td colspan="2">Header</td></tr><tr><td>A</td><td>B</td></tr></table>'
        b.attrs['raw_html'] = '<table><tr><td>Header</td><td></td></tr><tr><td>A</td><td>B</td></tr></table>'

        # 结构不同，TEDS 不应为 1.0
        score = compute_table_teds(a, b)
        self.assertLess(score, 0.95)
        self.assertGreater(score, 0.0)

        # 空表格评估应返回 0.0 而非 1.0
        self.assertEqual(compute_table_teds("", ""), 0.0)
        self.assertEqual(compute_table_teds(pd.DataFrame(), pd.DataFrame()), 0.0)
        self.assertEqual(compute_table_teds(None, None), 0.0)

    def test_load_excel_with_attrs_preserves_strings_and_formulas(self):
        """P2-6: _load_excel_with_attrs 重读旧工作簿时，编号 '001' 不变 1、'NA' 保留、公式不丢失"""
        from openpyxl import Workbook
        import excel_export as ex
        with tempfile.TemporaryDirectory() as td:
            p = os.path.join(td, 'test_preserve.xlsx')
            wb = Workbook()
            ws = wb.active
            ws.append(['Sample', 'Value'])
            ws.append(['001', 2.5])
            ws.append(['NA', '=SUM(1,2)'])
            wb.save(p)

            loaded = ex._load_excel_with_attrs(p)
            self.assertEqual(loaded.iloc[0, 0], '001')
            self.assertEqual(loaded.iloc[1, 0], 'NA')
            self.assertEqual(loaded.iloc[1, 1], '=SUM(1,2)')


if __name__ == '__main__':
    unittest.main()
