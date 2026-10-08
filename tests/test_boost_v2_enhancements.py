#!/usr/bin/env python3
"""
test_boost_v2_enhancements.py — 针对全新集成提升特性的深度单元测试套件：
1. arXiv HTML (ar5iv / arXiv LaTeXML) & NCBI PMC (PubMed Central) Open Access JATS XML 100% 原生直取策略
2. openpyxl CellRichText 学术上下标 (10^{-3}, CaCO_3, H_2O, Fe^{3+}) 与化学/数学格式保真导出
3. Artifex / PyMuPDF4LLM 原生 PDF 表格提取与 8 引擎多方投票池融合
4. IBM TEDS (Tree-Edit-Distance-based Similarity) 表格拓扑结构与内容度量算法
5. AllenAI olmOCR 2 学术表格边界与尾注分离启发式 (* p < 0.05, Note: ..., 注: ...)
"""

import os
import re
import sys
import tempfile
import unittest
from unittest.mock import patch, MagicMock
from pathlib import Path

import pandas as pd
import numpy as np

# 确保 scripts 目录在 sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

from scripts import common
from scripts import teds
from scripts import models
from scripts import excel_export
from scripts import table_postprocess
from scripts import pdf_table_extractor
from scripts.online import doi_resolver, strategies, graph


class TestArxivAndNCBIPMCStrategies(unittest.TestCase):
    """测试 1: arXiv HTML 与 NCBI PMC JATS XML 高保真原生表格抓取策略。"""

    def test_arxiv_id_extraction(self):
        """测试从 DOI、URL、字符串中精准提取各种 arXiv 编号。"""
        self.assertEqual(doi_resolver.extract_arxiv_id("10.48550/arXiv.2301.12345"), "2301.12345")
        self.assertEqual(doi_resolver.extract_arxiv_id("10.48550/arXiv.2301.12345v2"), "2301.12345v2")
        self.assertEqual(doi_resolver.extract_arxiv_id("https://arxiv.org/abs/2402.09999"), "2402.09999")
        self.assertEqual(doi_resolver.extract_arxiv_id("https://arxiv.org/html/2402.09999v1"), "2402.09999v1")
        self.assertEqual(doi_resolver.extract_arxiv_id("https://ar5iv.labs.arxiv.org/html/2205.12345"), "2205.12345")
        self.assertEqual(doi_resolver.extract_arxiv_id("arXiv:2106.01234"), "2106.01234")
        self.assertEqual(doi_resolver.extract_arxiv_id("math/0211159"), "math/0211159")
        self.assertIsNone(doi_resolver.extract_arxiv_id("10.1016/j.jvolgeores.2020.106888"))

    def test_pmcid_extraction_and_resolution(self):
        """测试 PMCID 提取及通过 NCBI API 从 DOI 解析 PMCID。"""
        self.assertEqual(doi_resolver.extract_pmcid("PMC8123456"), "PMC8123456")
        self.assertEqual(doi_resolver.extract_pmcid("https://pmc.ncbi.nlm.nih.gov/articles/PMC7654321/"), "PMC7654321")

        # Mock NCBI ID converter API
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.json.return_value = {
            "status": "ok",
            "records": [{"doi": "10.1371/journal.pone.0251234", "pmcid": "PMC8123456"}]
        }
        with patch("requests.get", return_value=mock_resp):
            resolved = doi_resolver.resolve_pmcid_from_doi("10.1371/journal.pone.0251234")
            self.assertEqual(resolved, "PMC8123456")

    def test_parse_jats_xml_tables(self):
        """测试解析标准 NISO JATS XML <table-wrap> 表格（含 colspan, rowspan, caption, 尾注）。"""
        sample_jats = """<?xml version="1.0"?>
        <article>
          <body>
            <table-wrap id="tbl1">
              <label>Table 1</label>
              <caption><title>Baseline clinical characteristics</title></caption>
              <table frame="hsides" rules="groups">
                <thead>
                  <tr>
                    <th align="left">Variable</th>
                    <th align="center" colspan="2">Cohort Distribution</th>
                  </tr>
                  <tr>
                    <th></th>
                    <th>Group A (n=40)</th>
                    <th>Group B (n=40)</th>
                  </tr>
                </thead>
                <tbody>
                  <tr>
                    <td>Age (years)</td>
                    <td>42.5</td>
                    <td>44.1</td>
                  </tr>
                  <tr>
                    <td>BMI</td>
                    <td>23.1</td>
                    <td>24.0</td>
                  </tr>
                </tbody>
              </table>
              <table-wrap-foot>
                <fn id="fn1"><p>* p &lt; 0.05 statistically significant.</p></fn>
              </table-wrap-foot>
            </table-wrap>
          </body>
        </article>
        """
        tables = strategies._parse_jats_xml_tables(sample_jats)
        self.assertEqual(len(tables), 1)
        df = tables[0]
        self.assertEqual(df.attrs.get("label"), "Table 1")
        self.assertTrue("Baseline clinical characteristics" in df.attrs.get("table_title", ""))
        self.assertEqual(df.attrs.get("extractor"), "ncbi_pmc_jats")
        self.assertEqual(df.attrs.get("footnotes"), ["* p < 0.05 statistically significant."])
        self.assertEqual(df.shape[1], 3)
        self.assertTrue(any("Age" in str(v) for v in df.values.flatten()))

    def test_extract_tables_via_ncbi_pmc_mock(self):
        """测试 extract_tables_via_ncbi_pmc 端到端抓取与调用。"""
        sample_jats = """
        <table-wrap id="t1">
          <label>Table 1</label>
          <caption><p>Sample measurements</p></caption>
          <table>
            <tr><th>Sample</th><th>Value</th></tr>
            <tr><td>S1</td><td>99.5</td></tr>
          </table>
        </table-wrap>
        """
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.text = sample_jats
        with patch("requests.get", return_value=mock_resp):
            dfs = strategies.extract_tables_via_ncbi_pmc("PMC1234567")
            self.assertEqual(len(dfs), 1)
            self.assertEqual(dfs[0].attrs.get("label"), "Table 1")

    def test_extract_tables_via_arxiv_html_mock(self):
        """测试 extract_tables_via_arxiv_html 解析 arXiv 官方 HTML / ar5iv LaTeXML 表格。"""
        sample_arxiv_html = """
        <!DOCTYPE html>
        <html>
        <body>
          <figure class="ltx_table" id="S1.T1">
            <figcaption class="ltx_caption">
              <span class="ltx_tag ltx_tag_table">Table 1: </span>Benchmark evaluation results on GLUE
            </figcaption>
            <table class="ltx_tabular">
              <tr>
                <th class="ltx_th">Model</th>
                <th class="ltx_th">Accuracy</th>
                <th class="ltx_th">F1 Score</th>
              </tr>
              <tr>
                <td class="ltx_td">Baseline</td>
                <td class="ltx_td">85.4</td>
                <td class="ltx_td">84.9</td>
              </tr>
            </table>
            <div class="ltx_table_notes">* Higher is better.</div>
          </figure>
        </body>
        </html>
        """
        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.text = sample_arxiv_html
        with patch("requests.get", return_value=mock_resp):
            dfs = strategies.extract_tables_via_arxiv_html("2401.12345")
            self.assertEqual(len(dfs), 1)
            df = dfs[0]
            self.assertEqual(df.attrs.get("label"), "Table 1")
            self.assertTrue("Benchmark evaluation" in df.attrs.get("table_title", ""))
            self.assertEqual(df.attrs.get("footnotes"), ["* Higher is better."])
            self.assertEqual(df.attrs.get("extractor"), "arxiv_html")

    def test_graph_nodes_arxiv_and_pmc(self):
        """测试 ArxivHTMLNode 与 NCBIPMCNode 在 DAG 与分层竞速中的调度。"""
        node_arxiv = graph.ArxivHTMLNode()
        node_pmc = graph.NCBIPMCNode()

        dummy_df = pd.DataFrame({"A": [1, 2], "B": [3, 4]})
        with patch("scripts.online.graph.extract_tables_via_arxiv_html", return_value=[dummy_df]):
            state_arxiv = {"url": "https://arxiv.org/abs/2301.00001"}
            res_state, next_node = node_arxiv.execute(state_arxiv)
            self.assertIsNone(next_node)
            self.assertEqual(len(res_state["dfs"]), 1)

        with patch("scripts.online.graph.extract_tables_via_ncbi_pmc", return_value=[dummy_df]):
            state_pmc = {"doi": "10.1371/journal.pone.0001", "pmcid": "PMC1234567"}
            res_state, next_node = node_pmc.execute(state_pmc)
            self.assertIsNone(next_node)
            self.assertEqual(len(res_state["dfs"]), 1)


class TestOpenpyxlCellRichTextAcademicFormatting(unittest.TestCase):
    """测试 2: openpyxl CellRichText 原生学术上下标与化学/数学保真导出。"""

    def test_parse_academic_runs(self):
        """测试学术上下标文本分段解析 (LaTeX 与 Unicode)。"""
        # LaTeX 上标
        runs_sup = table_postprocess.parse_academic_runs("10^{-3}")
        self.assertEqual(runs_sup, [("10", None), ("-3", "superscript")])

        # 化学分子式下标
        runs_chem = table_postprocess.parse_academic_runs("CaCO_3")
        self.assertEqual(runs_chem, [("CaCO", None), ("3", "subscript")])

        # 离子价态
        runs_ion = table_postprocess.parse_academic_runs("Fe^{3+}")
        self.assertEqual(runs_ion, [("Fe", None), ("3+", "superscript")])

        # Unicode 上下标
        runs_uni = table_postprocess.parse_academic_runs("10⁻³ mol/L")
        self.assertEqual(runs_uni, [("10", None), ("-3", "superscript"), (" mol/L", None)])

        runs_h2o = table_postprocess.parse_academic_runs("H₂O")
        self.assertEqual(runs_h2o, [("H", None), ("2", "subscript"), ("O", None)])

        # LaTeX H_2O 与 Al_2O_3 (防止元素符号被吞入下标)
        runs_h2o_latex = table_postprocess.parse_academic_runs("H_2O")
        self.assertEqual(runs_h2o_latex, [("H", None), ("2", "subscript"), ("O", None)])

        runs_al2o3 = table_postprocess.parse_academic_runs("Al_2O_3")
        self.assertEqual(runs_al2o3, [("Al", None), ("2", "subscript"), ("O", None), ("3", "subscript")])

        # 保护 snake_case 变量标识符 (sample_id, table_1, step_1, group_2) 绝不误杀为下标
        self.assertEqual(table_postprocess.parse_academic_runs("sample_id"), [("sample_id", None)])
        self.assertEqual(table_postprocess.parse_academic_runs("table_1"), [("table_1", None)])
        self.assertEqual(table_postprocess.parse_academic_runs("group_2"), [("group_2", None)])
        self.assertEqual(table_postprocess.parse_academic_runs("step_1"), [("step_1", None)])

    def test_format_academic_rich_text_openpyxl(self):
        """测试转换为 openpyxl CellRichText 对象。"""
        from openpyxl.cell.rich_text import CellRichText

        # 含上下标字符成功转为 CellRichText
        rt1 = excel_export.format_academic_rich_text("10^{-3}")
        self.assertIsInstance(rt1, CellRichText)
        self.assertEqual(str(rt1), "10-3")

        rt2 = excel_export.format_academic_rich_text("CaCO_3")
        self.assertIsInstance(rt2, CellRichText)
        self.assertEqual(str(rt2), "CaCO3")

        # 严格保留计算用数值型 (float, int)，防止破坏公式
        num_val = 123.456
        self.assertEqual(excel_export.format_academic_rich_text(num_val), num_val)
        self.assertIsInstance(excel_export.format_academic_rich_text(num_val), float)

        num_str = "45.67"
        self.assertEqual(excel_export.format_academic_rich_text(num_str), "45.67")

    def test_excel_export_with_rich_text_roundtrip(self):
        """测试将含上下标的数据表导出为 Excel，并验证文件完整性。"""
        import openpyxl

        df = pd.DataFrame({
            "Sample": ["Quartz", "Calcite"],
            "Formula": ["SiO_2", "CaCO_3"],
            "Concentration": ["10^{-4} mol/L", "10^{-2} mol/L"],
            "NumericValue": [12.5, 34.8]
        })

        with tempfile.TemporaryDirectory() as tmp_dir:
            out_file = os.path.join(tmp_dir, "Table 1.xlsx")
            success = excel_export.save_tables_to_excel([df], out_file, single_file=True)
            self.assertTrue(success)
            self.assertTrue(os.path.exists(out_file))

            # 重新加载验证
            wb = openpyxl.load_workbook(out_file)
            ws = wb.active
            # 验证纯数字单元格保持 float 适合计算
            self.assertIsInstance(ws.cell(row=2, column=4).value, (float, int))
            # 验证富文本单元格渲染且可读取字符串内容
            f_val = ws.cell(row=3, column=2).value
            self.assertTrue("CaCO3" in str(f_val))


class TestPyMuPDF4LLMIntegration(unittest.TestCase):
    """测试 3: Artifex PyMuPDF4LLM 候选提取器与 8 引擎多方投票。"""

    def test_pymupdf4llm_availability_probe(self):
        """测试 is_pymupdf4llm_available 正确探测模块状态。"""
        is_avail = pdf_table_extractor.is_pymupdf4llm_available()
        self.assertIsInstance(is_avail, bool)

    def test_extract_via_pymupdf4llm_mock(self):
        """测试 extract_via_pymupdf4llm 使用 Mock 解析返回标准化表格列表。"""
        mock_pymupdf4llm = MagicMock()
        mock_pymupdf4llm.to_markdown.return_value = """
# Section 1
Here is the table:

| Element | Weight% | Error |
|---|---|---|
| SiO2 | 52.3 | 0.2 |
| FeO | 8.1 | 0.1 |

End of section.
"""
        with tempfile.NamedTemporaryFile(suffix=".pdf") as tmp_pdf:
            import fitz
            doc = fitz.open()
            doc.new_page()
            doc.save(tmp_pdf.name)
            doc.close()
            with patch("scripts.pdf_table_extractor.is_pymupdf4llm_available", return_value=True), \
                 patch.dict("sys.modules", {"pymupdf4llm": mock_pymupdf4llm}):
                results = pdf_table_extractor.extract_via_pymupdf4llm(tmp_pdf.name, pages=[0])

        self.assertEqual(len(results), 1)
        res = results[0]
        self.assertEqual(res["page_idx"], 0)
        self.assertEqual(res["df"].attrs.get("extractor"), "pymupdf4llm")
        self.assertEqual(res["df"].shape, (2, 3))
        self.assertEqual(res["df"].iloc[0, 0], "SiO2")

    def test_three_way_vote_with_pymupdf4llm(self):
        """测试 8 引擎投票池正确接纳 pymupdf4llm_tables 并执行投票。"""
        df_target = pd.DataFrame({"Col1": ["A", "B"], "Col2": [1, 2]})
        df_target.attrs["extractor"] = "pymupdf4llm"

        inspector_tables = []
        ft_tables = []
        plumber_tables = [{"df": df_target.copy(), "page_idx": 0, "table_idx": 0}]
        camelot_tables = []
        pymupdf_tables = [{"df": df_target.copy(), "page_idx": 0, "table_idx": 0}]

        fused, logs, summary = pdf_table_extractor.three_way_vote(
            inspector_tables, ft_tables, plumber_tables, camelot_tables,
            pymupdf4llm_tables=pymupdf_tables
        )
        self.assertEqual(len(fused), 1)
        # 优先级排序中 pymupdf4llm 优先于 pdfplumber
        self.assertEqual(fused[0]["df"].attrs.get("extractor"), "pymupdf4llm")


class TestIBMTEDSTableTopologicalMetric(unittest.TestCase):
    """测试 4: IBM TEDS (Tree-Edit-Distance-based Similarity) 表格拓扑与结构度量。"""

    def test_dataframe_to_table_tree(self):
        """测试 DataFrame 准确映射为 DOM 树结构。"""
        df = pd.DataFrame({"A": [10, 20], "B": [30, 40]})
        tree = teds.dataframe_to_table_tree(df)
        self.assertEqual(tree.tag, "table")
        # 1 行表头 + 2 行数据 = 3 个 tr
        self.assertEqual(len(tree.children), 3)
        self.assertEqual([c.tag for c in tree.children[0].children], ["th", "th"])
        self.assertEqual([c.tag for c in tree.children[1].children], ["td", "td"])

    def test_parse_html_to_tree_with_spans(self):
        """测试 HTML table 准确解析并保留 colspan 与 rowspan。"""
        html_tbl = """
        <table>
          <tr>
            <th colspan="2">Merged Header</th>
          </tr>
          <tr>
            <td rowspan="2">Span Row</td>
            <td>Cell 1</td>
          </tr>
          <tr>
            <td>Cell 2</td>
          </tr>
        </table>
        """
        tree = teds.parse_html_to_tree(html_tbl)
        self.assertEqual(len(tree.children), 3)
        th_cell = tree.children[0].children[0]
        self.assertEqual(th_cell.attrs.get("colspan"), 2)
        td_cell = tree.children[1].children[0]
        self.assertEqual(td_cell.attrs.get("rowspan"), 2)

    def test_teds_similarity_exact_and_variations(self):
        """测试 TEDS 相似度计算数值边界与属性敏感性。"""
        df1 = pd.DataFrame({"Element": ["Si", "Al", "Fe"], "Val": [50.1, 15.2, 8.4]})
        df2 = pd.DataFrame({"Element": ["Si", "Al", "Fe"], "Val": [50.1, 15.2, 8.4]})
        # 完全一致返回 1.0
        self.assertAlmostEqual(teds.teds_similarity(df1, df2), 1.0, places=4)

        # 单元格数值轻微变化，相似度高但略低于 1.0
        df3 = pd.DataFrame({"Element": ["Si", "Al", "Fe"], "Val": [50.1, 15.2, 9.9]})
        sim_diff = teds.teds_similarity(df1, df3)
        self.assertTrue(0.85 < sim_diff < 1.0)

        # 结构一致但内容不同时，structure_only=True 应返回 1.0
        self.assertAlmostEqual(teds.teds_similarity(df1, df3, structure_only=True), 1.0, places=4)

        # 完全不同维度的表格
        df_unrelated = pd.DataFrame({"City": ["Beijing", "Shanghai", "Guangzhou", "Shenzhen"], "Pop": [21, 24, 18, 17], "GDP": [4, 4.5, 3, 3.2]})
        sim_unrelated = teds.teds_similarity(df1, df_unrelated)
        self.assertTrue(sim_unrelated < 0.6)

    def test_table_similarity_integration(self):
        """测试 pdf_table_extractor._table_similarity 结合 TEDS 拓扑度量后的综合准确性。"""
        df1 = pd.DataFrame({"A": [1, 2], "B": [3, 4]})
        df2 = pd.DataFrame({"A": [1, 2], "B": [3, 4]})
        self.assertAlmostEqual(pdf_table_extractor._table_similarity(df1, df2), 1.0, places=4)


class TestOlmOCR2BoundaryAndFootnoteHeuristics(unittest.TestCase):
    """测试 5: AllenAI olmOCR 2 学术表格边界与尾注分离启发式。"""

    def test_is_academic_footnote_row(self):
        """测试对各种学术注释行模式的识别。"""
        # 统计显著性
        is_note, text = table_postprocess.is_academic_footnote_row(["* p < 0.05 compared with control group.", ""], 2)
        self.assertTrue(is_note)
        self.assertTrue("p < 0.05" in text)

        # Note / 注明前缀
        is_note, text = table_postprocess.is_academic_footnote_row(["Note: All measurements in triplicate.", "triplicate"], 2)
        self.assertTrue(is_note)
        self.assertTrue("All measurements" in text)

        # 中文注：
        is_note, text = table_postprocess.is_academic_footnote_row(["注：数据来源于中国统计年鉴（2022年）。", ""], 2)
        self.assertTrue(is_note)

        # 数据来源 Source:
        is_note, text = table_postprocess.is_academic_footnote_row(["Source: USGS Mineral Commodity Summaries (2023).", ""], 2)
        self.assertTrue(is_note)

        # 正常数据行绝不误杀为尾注
        is_data, _ = table_postprocess.is_academic_footnote_row(["Sample A", "12.5", "3.2"], 3)
        self.assertFalse(is_data)

    def test_separate_table_footnotes_clean_dataframe(self):
        """测试 separate_table_footnotes 剥离尾部注释并存入 df.attrs['footnotes']。"""
        df = pd.DataFrame({
            "Sample": ["S1", "S2", "* p < 0.05 vs baseline.", "Note: BDL = Below Detection Limit."],
            "Wt%": ["54.2", "48.9", "", ""]
        })

        clean_df = table_postprocess.separate_table_footnotes(df)
        self.assertEqual(len(clean_df), 2)
        self.assertEqual(clean_df.iloc[-1, 0], "S2")

        footnotes = clean_df.attrs.get("footnotes", [])
        self.assertEqual(len(footnotes), 2)
        self.assertTrue(any("p < 0.05" in fn for fn in footnotes))
        self.assertTrue(any("BDL" in fn for fn in footnotes))

    def test_separate_table_boundary_caption_in_row0(self):
        """测试第 0 行误将 Table Caption 纳入数据体时的边界隔离与表题提升。"""
        df = pd.DataFrame({
            "Col1": ["Table 2. Summary of experimental results", "Sample 1", "Sample 2"],
            "Col2": ["", "100", "200"]
        })

        clean_df = table_postprocess.separate_table_footnotes(df)
        self.assertEqual(len(clean_df), 2)
        self.assertEqual(clean_df.iloc[0, 0], "Sample 1")
        self.assertEqual(clean_df.attrs.get("label"), "Table 2")
        self.assertEqual(clean_df.attrs.get("table_title"), "Table 2. Summary of experimental results")

    def test_postprocess_dataframe_integrates_footnote_separation(self):
        """测试 postprocess_dataframe 流水线自动完成尾注剥离。"""
        df = pd.DataFrame({
            "Group": ["Control", "Treated", "Source: Experimental data."],
            "Score": ["10.2", "14.8", ""]
        })

        processed = table_postprocess.postprocess_dataframe(df)
        self.assertEqual(len(processed), 2)
        self.assertTrue("footnotes" in processed.attrs)
        self.assertTrue("Source: Experimental data." in processed.attrs["footnotes"][0])

    def test_extracted_table_syncs_footnotes(self):
        """测试 ExtractedTable 统一对象中 footnotes 属性的双向同步。"""
        df = pd.DataFrame({"A": [1, 2], "B": [3, 4]})
        et = models.ExtractedTable(
            df=df,
            label="Table 1",
            footnotes=["* p < 0.05"]
        )
        self.assertEqual(et.footnotes, ["* p < 0.05"])
        self.assertEqual(et.to_dataframe().attrs.get("footnotes"), ["* p < 0.05"])

    def test_clean_vlm_dataframe_footnotes(self):
        """测试 _clean_vlm_dataframe 针对 VLM 提取结果剥离尾注并保留 df.attrs['footnotes']。"""
        df = pd.DataFrame({
            "Mineral": ["Quartz", "Pyrite", "* p < 0.05 vs baseline"],
            "Content": ["80%", "20%", ""]
        })
        orig_attrs = {"source": "vlm", "page_idx": 1}
        cleaned = table_postprocess._clean_vlm_dataframe(df, orig_attrs)
        self.assertEqual(len(cleaned), 2)
        self.assertEqual(cleaned.iloc[-1, 0], "Pyrite")
        self.assertIn("footnotes", cleaned.attrs)
        self.assertEqual(cleaned.attrs["footnotes"], ["* p < 0.05 vs baseline"])
        self.assertEqual(cleaned.attrs["source"], "vlm")

    def test_formula_preservation_in_excel(self):
        """测试 Excel 公式单元格 (=SUM(...)) 绝对保持公式数据类型，不被转为 CellRichText。"""
        formula = "=SUM(A1:A10)"
        res = excel_export.format_academic_rich_text(formula)
        self.assertEqual(res, formula)
        self.assertIsInstance(res, str)

        import openpyxl
        with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws["A1"] = 10
            ws["A2"] = 20
            ws["A3"] = "=SUM(A1:A2)"
            ws["B1"] = "CaCO_3"
            wb.save(tmp.name)

            excel_export.autofit_excel_columns(tmp.name)

            wb2 = openpyxl.load_workbook(tmp.name)
            ws2 = wb2.active
            # 验证公式单元格仍为公式且以 = 开头
            self.assertEqual(ws2["A3"].value, "=SUM(A1:A2)")
            # 验证化学分子式转换为富文本
            self.assertTrue("CaCO3" in str(ws2["B1"].value))

    def test_parse_jats_xml_with_namespaces(self):
        """测试解析含 XML 命名空间 (如 xmlns="http://www.w3.org/1999/xhtml") 的真实 JATS XML。"""
        jats_with_ns = """<?xml version="1.0"?>
        <article xmlns:xlink="http://www.w3.org/1999/xlink">
          <body>
            <table-wrap id="tbl2">
              <label>Table 2</label>
              <caption><title>Experimental parameters</title></caption>
              <table xmlns="http://www.w3.org/1999/xhtml">
                <thead>
                  <tr><th>Run</th><th>Temp</th><th>Yield</th></tr>
                </thead>
                <tbody>
                  <tr><td>1</td><td>350 K</td><td>88.2%</td></tr>
                  <tr><td>2</td><td>375 K</td><td>93.5%</td></tr>
                </tbody>
              </table>
              <table-wrap-foot>
                <p>Note: Yield determined by GC-MS.</p>
              </table-wrap-foot>
            </table-wrap>
          </body>
        </article>
        """
        dfs = strategies._parse_jats_xml_tables(jats_with_ns)
        self.assertEqual(len(dfs), 1)
        df = dfs[0]
        self.assertEqual(df.attrs.get("label"), "Table 2")
        self.assertTrue("Experimental parameters" in df.attrs.get("table_title", ""))
        self.assertEqual(df.shape, (2, 3))
        self.assertEqual(df.attrs.get("footnotes"), ["Note: Yield determined by GC-MS."])

    def test_config_toggles_respected(self):
        """测试 TEDS_METRIC_ENABLED, ARXIV_HTML_ENABLED, NCBI_PMC_ENABLED 配置开关生效。"""
        # 1. TEDS 禁用
        with patch("common.load_config", return_value={"TEDS_METRIC_ENABLED": False}), \
             patch("scripts.pdf_table_extractor.teds_similarity") as mock_teds:
            df1 = pd.DataFrame({"A": [1, 2]})
            df2 = pd.DataFrame({"A": [1, 2]})
            sim = pdf_table_extractor._table_similarity(df1, df2)
            self.assertGreaterEqual(sim, 0.9)
            mock_teds.assert_not_called()

        # 2. arXiv HTML 禁用
        with patch("scripts.online.strategies.load_config", return_value={"ARXIV_HTML_ENABLED": False}):
            res = strategies.extract_tables_via_arxiv_html("2401.12345")
            self.assertEqual(res, [])

        # 3. NCBI PMC 禁用
        with patch("scripts.online.strategies.load_config", return_value={"NCBI_PMC_ENABLED": False}):
            res = strategies.extract_tables_via_ncbi_pmc("PMC1234567")
            self.assertEqual(res, [])

    def test_teds_wide_and_tall_table_performance(self):
        """测试 TEDS 拓扑算法在超大宽表（100行 x 50列）下的快速自适应采样与毫秒级执行。"""
        import time
        df_large1 = pd.DataFrame([[f"val_{r}_{c}" for c in range(50)] for r in range(100)])
        df_large2 = pd.DataFrame([[f"val_{r}_{c}" for c in range(50)] for r in range(100)])

        t0 = time.time()
        sim = teds.teds_similarity(df_large1, df_large2, max_sample_rows=15, max_sample_cols=10)
        t_elapsed = time.time() - t0

        self.assertAlmostEqual(sim, 1.0, places=3)
        self.assertLess(t_elapsed, 0.3)  # 应在 300ms 甚至更短内完成


if __name__ == "__main__":
    unittest.main()
