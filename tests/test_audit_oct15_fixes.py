#!/usr/bin/env python3
"""
tests/test_audit_oct15_fixes.py
Comprehensive regression tests for the 9 issues identified in the Oct 15 audit.
"""

import os
import sys
import tempfile
import unittest
import types
import io
import contextlib
from pathlib import Path
from unittest.mock import patch, MagicMock

import pandas as pd
import numpy as np
import pymupdf as fitz
from PIL import Image, ImageDraw
from openpyxl import load_workbook

# Add scripts directory to path
SCRIPTS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "scripts")
if SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, SCRIPTS_DIR)

import common
import ocr_client as oc
import pdf_table_extractor as pte
import doclayout_yolo_detector as dl
import excel_export as ex
import table_validator as tv
from online.strategies import _parse_jats_xml_tables, _rows_to_dataframe, _xml_cell_to_text, _bs4_cell_to_text
from table_postprocess import parse_academic_runs
import rapid_table_extractor as rte


class TestAuditOct15Fixes(unittest.TestCase):

    def test_issue1_multipage_rescue_tracking_does_not_starve_remaining_pages(self):
        """[P1] 一页补救成功，不能阻断其他页面继续补漏。"""
        with tempfile.TemporaryDirectory() as td:
            pdf_path = str(Path(td) / "twopage.pdf")
            doc = fitz.open()
            for _ in range(2):
                p = doc.new_page()
                p.insert_text((50, 50), "Searchable narrative text for testing.")
            doc.save(pdf_path)
            doc.close()

            # Page 0 has a rescued table from rapid_table
            rapid_df = pd.DataFrame({"A": [1, 2], "B": [3, 4]})
            rapid_df.attrs.update(page_idx=0, extractor="rapid_table")
            rapid_item = {'df': rapid_df, 'page_idx': 0, 'bbox': [10, 10, 100, 100]}

            # Page 1 has a table rescued from paddleocr_crops
            crop_df = pd.DataFrame({"C": [5, 6], "D": [7, 8]})
            crop_df.attrs.update(page_idx=1, extractor="paddleocr_vl_crop")
            crop_item = {'df': crop_df, 'page_idx': 1, 'bbox': [20, 20, 200, 200]}

            inspector = {'pdf_type': 'text_based', 'confidence': 1.0, 'pages_with_tables': [1, 2], 'tables': [], 'error': None, 'pages_needing_ocr': []}

            with contextlib.ExitStack() as stack:
                stack.enter_context(patch.object(tv, 'scan_pdf_table_declarations', return_value={}))
                stack.enter_context(patch.object(tv, 'validate_native_extraction_pipeline', return_value=([], {0, 1}, [])))
                stack.enter_context(patch.object(pte, 'classify_and_extract_via_inspector', return_value=inspector))
                for name in ['extract_via_structure_tree', 'extract_via_find_tables', 'extract_via_pdfplumber', 'extract_via_camelot', 'extract_via_img2table', 'extract_via_gmft', 'extract_via_pymupdf4llm', 'extract_via_text_alignment']:
                    stack.enter_context(patch.object(pte, name, return_value=[]))
                stack.enter_context(patch.object(pte, 'is_rapid_table_available', return_value=True))
                # Rapid returns table for page 0 only
                stack.enter_context(patch.object(pte, 'extract_via_rapid_table', return_value=[rapid_item]))
                stack.enter_context(patch.object(pte, 'three_way_vote', return_value=([], [], {})))
                stack.enter_context(patch.object(pte, 'is_table_low_quality', return_value=False))
                stack.enter_context(patch.object(pte, 'reconcile_table_captions', side_effect=lambda r, p: r))
                crops_mock = stack.enter_context(patch.object(pte, 'extract_via_paddleocr_crops', return_value=[crop_item]))

                results, logs = pte._do_extract_tables_from_pdf(pdf_path)

                # Page 0 rescue must NOT starve page 1 from running crops
                crops_mock.assert_called_once_with(pdf_path, [1])
                result_pages = sorted(r['page_idx'] for r in results)
                self.assertEqual(result_pages, [0, 1])

    def test_issue2_online_numeric_rows_not_dropped(self):
        """[P1] 在线解析不能仅凭数字比例误删合法的纯数值首行数据。"""
        # Test 1: Two numeric rows in JATS table must keep both rows
        numeric_xml = '<article><table-wrap><label>Table 1</label><table><tbody><tr><td>10</td><td>20</td></tr><tr><td>30</td><td>40</td></tr></tbody></table></table-wrap></article>'
        jats_tables = _parse_jats_xml_tables(numeric_xml)
        self.assertEqual(len(jats_tables), 1)
        self.assertEqual(jats_tables[0].shape, (2, 2))
        self.assertEqual(jats_tables[0].values.tolist(), [['10', '20'], ['30', '40']])

        # Test 2: Explicit rows_to_dataframe call with numeric data
        raw_numeric = [['10', '20'], ['30', '40']]
        df = _rows_to_dataframe(raw_numeric)
        self.assertEqual(df.shape, (2, 2))
        self.assertEqual(df.values.tolist(), [['10', '20'], ['30', '40']])

        # Test 3: True index row [0, 1] followed by text header [ColA, ColB] should still drop index row
        indexed_data = [['0', '1'], ['ColA', 'ColB'], ['v1', 'v2']]
        df_idx = _rows_to_dataframe(indexed_data)
        self.assertEqual(list(df_idx.columns), ['ColA', 'ColB'])
        self.assertEqual(df_idx.values.tolist(), [['v1', 'v2']])

    def test_issue3_gmft_bbox_and_zero_rect_handling(self):
        """[P1] GMFT 表格坐标读取支持 rect.bbox 与 xmin/ymin/xmax/ymax，且过滤全零框。"""
        with tempfile.TemporaryDirectory() as td:
            pdf_path = str(Path(td) / "gmft.pdf")
            doc = fitz.open()
            doc.new_page().insert_text((30, 30), "Some text")
            doc.save(pdf_path)
            doc.close()

            class MockDoc:
                def __init__(self, *a): pass
                def __len__(self): return 1
                def get_page(self, i): return i
                def close(self): pass

            # Table with official rect.bbox / xmin/ymin/xmax/ymax
            valid_table = types.SimpleNamespace(
                rect=types.SimpleNamespace(bbox=(10, 20, 200, 300), xmin=10, ymin=20, xmax=200, ymax=300)
            )
            # Table with invalid zero-area rect [0, 0, 0, 0]
            zero_table = types.SimpleNamespace(
                rect=types.SimpleNamespace(bbox=(0, 0, 0, 0), xmin=0, ymin=0, xmax=0, ymax=0)
            )

            mock_detector = types.SimpleNamespace(extract=lambda page: [valid_table, zero_table])
            mock_formatter = types.SimpleNamespace(
                extract=lambda t: types.SimpleNamespace(df=lambda: pd.DataFrame({'A': [1], 'B': [2]}))
            )

            mods = {
                'gmft.auto': types.SimpleNamespace(
                    AutoTableDetector=lambda: mock_detector,
                    AutoTableFormatter=lambda: mock_formatter
                ),
                'gmft.pdf_bindings': types.SimpleNamespace(PyPDFium2Document=MockDoc)
            }

            with patch.dict(sys.modules, mods), patch.object(common, 'load_config', return_value={}), patch.object(pte, 'is_gmft_available', return_value=True):
                res = pte.extract_via_gmft(pdf_path, pages=[0])
                self.assertEqual(len(res), 2)
                # Valid table gets [10, 20, 200, 300]
                self.assertEqual(res[0]['bbox'], [10, 20, 200, 300])
                # Zero-area table has bbox filtered out to None
                self.assertIsNone(res[1]['bbox'])

    def test_issue4_pp_doclayout_coords_no_double_scaling(self):
        """[P1] PP-DocLayoutV3 在包含 scale_factor 输入时不重复缩放原图坐标。"""
        class SessionWithScale:
            def get_inputs(self):
                return [
                    types.SimpleNamespace(name='image', shape=[1, 3, 800, 800]),
                    types.SimpleNamespace(name='scale_factor', shape=[1, 2])
                ]
            def run(self, *args):
                # Class 21 (table), score 0.95, coords already in image space [100, 100, 500, 300]
                return [np.array([[21, 0.95, 100, 100, 500, 300]], dtype=float)]

        detector = dl.PPDocLayoutV3Detector.__new__(dl.PPDocLayoutV3Detector)
        detector._available = True
        detector._session = SessionWithScale()

        # Input image is 1600x800
        img = Image.new("RGB", (1600, 800), "white")
        dets = detector.detect(img)
        self.assertEqual(len(dets), 1)
        self.assertEqual(dets[0]['bbox'], [100.0, 100.0, 500.0, 300.0])

    def test_issue5_jats_and_arxiv_rich_text_fidelity(self):
        """[P1] 在线原生表格到 Excel 的保真链路：保留上下标、尾注与 openpyxl CellRichText。"""
        xml = (
            '<article><table-wrap><label>Table 1</label>'
            '<table><thead><tr><th>Chemical</th><th>Conc</th></tr></thead>'
            '<tbody><tr><td>CO<sub>2</sub></td><td>10<sup>−3</sup></td></tr>'
            '<tr><td>H<sub>2</sub>O</td><td>5.5</td></tr></tbody></table>'
            '<table-wrap-foot><p><sup>a</sup> dry basis.</p></table-wrap-foot></table-wrap></article>'
        )
        tables = _parse_jats_xml_tables(xml)
        self.assertEqual(len(tables), 1)
        df = tables[0]
        # Footnote preserved
        self.assertTrue(any('dry basis' in fn for fn in df.attrs.get('footnotes', [])))

        # Export to excel
        with tempfile.TemporaryDirectory() as td:
            out_file = str(Path(td) / "table.xlsx")
            ex.save_tables_to_excel([df], out_file)
            wb = load_workbook(out_file, rich_text=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=False))

            # Verify footnote exists at the end
            all_str = " ".join(str(c.value) for row in rows for c in row if c.value is not None)
            self.assertIn("dry basis", all_str)

            # Test appending to existing file also runs autofit and rich text
            df2 = pd.DataFrame({"Species": ["CH<sub>4</sub>"], "Val": ["10<sup>−2</sup>"]})
            df2.attrs['label'] = "Table 2"
            df2.attrs['footnotes'] = ["Second table footnote"]
            ex.save_tables_to_excel([df2], out_file, single_file=True)

            wb2 = load_workbook(out_file, rich_text=True)
            self.assertIn("Table 2", wb2.sheetnames)
            ws2 = wb2["Table 2"]
            ws2_str = " ".join(str(c.value) for row in ws2.iter_rows(values_only=False) for c in row if c.value is not None)
            self.assertIn("Second table footnote", ws2_str)
            wb2.close()

    def test_issue6_docling_ocr_on_mixed_text_and_image_pages(self):
        """[P2] Docling 覆盖正文可搜索但内嵌栅格图像表格的页面。"""
        from docling.document_converter import PdfFormatOption
        from docling.datamodel.base_models import InputFormat

        with tempfile.TemporaryDirectory() as td:
            pdf_path = str(Path(td) / "mixed.pdf")
            doc = fitz.open()
            page = doc.new_page()
            # Narrative searchable text > 50 chars
            page.insert_text((30, 30), "This is sufficiently long searchable narrative body text exceeding fifty characters threshold.")
            # Insert a table image
            im = Image.new('RGB', (200, 100), 'white')
            buf = io.BytesIO()
            im.save(buf, format='PNG')
            page.insert_image(fitz.Rect(30, 70, 230, 170), stream=buf.getvalue())
            doc.save(pdf_path)
            doc.close()

            captured = {}
            class MockConverter:
                def __init__(self, format_options):
                    captured['ocr'] = format_options[InputFormat.PDF].pipeline_options.do_ocr
                def convert(self, p):
                    return types.SimpleNamespace(document=types.SimpleNamespace(tables=[]))

            with patch.object(oc, 'is_docling_available', return_value=True), patch.dict(
                sys.modules,
                {'docling.document_converter': types.SimpleNamespace(DocumentConverter=MockConverter, PdfFormatOption=PdfFormatOption)}
            ):
                oc.extract_via_docling(pdf_path, config={})
                self.assertTrue(captured.get('ocr'))

    def test_issue7_ocr_engine_dispatch_respects_configuration(self):
        """[P2] OCR 引擎选择完整生效，配置 OCR_ENGINE='docling' 时直接路由。"""
        with patch.object(oc, 'call_paddleocr_job') as cloud_job, patch.object(oc, 'call_docling_table_extractor', return_value="docling_md") as docling_job:
            # When OCR_ENGINE='docling'
            res = oc.call_paddleocr_vl_online_api("dummy.pdf", config={"OCR_ENGINE": "docling"})
            self.assertEqual(res, "docling_md")
            cloud_job.assert_not_called()
            docling_job.assert_called_once()

    def test_issue8_and_9_table_structure_rec_and_rapid_table_readiness(self):
        """[P2] RapidTable / TableStructureRec 官方组件接入与真实初始化状态检验。"""
        class MockTableCls:
            def __call__(self, img):
                return ('lineless', 0.02)

        class MockLinelessRec:
            def __call__(self, img):
                return ('<table><tr><th>Head</th></tr><tr><td>val</td></tr></table>', 0.05)

        mock_tsr = types.SimpleNamespace(
            TableCls=MockTableCls,
            LinelessTableRec=MockLinelessRec,
        )

        with patch.dict(sys.modules, {'table_structure_rec': mock_tsr}):
            rte._RAPID_TABLE_AVAILABLE = None
            self.assertTrue(rte.is_rapid_table_available())

            extractor = rte.RapidTableExtractor()
            self.assertTrue(extractor.is_ready())
            self.assertEqual(extractor.classify_table_type("dummy_img"), "wireless")
            df = extractor.extract_table_from_image("dummy_img")
            self.assertIsNotNone(df)
            self.assertEqual(df.values.tolist(), [['val']])


if __name__ == "__main__":
    unittest.main()
