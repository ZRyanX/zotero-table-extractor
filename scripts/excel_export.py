#!/usr/bin/env python3
"""
excel_export.py — 表格 Excel 导出（自 extract_zotero_table.py 拆出）。

包含：安全文件名、列宽自适应、多表 Excel 写出
（写出前过滤附表/元数据表，并合并跨页续表）。
"""

import os
import re

import pandas as pd

# Shared helpers (supplementary / metadata table filtering, label formatting)
try:
    from .common import is_supplementary, is_metadata_table, format_table_label, TABLE_LABEL_PATTERN, TABLE_LABEL_RE
except ImportError:
    from common import is_supplementary, is_metadata_table, format_table_label, TABLE_LABEL_PATTERN, TABLE_LABEL_RE

# Table post-processing helpers
try:
    from .table_postprocess import postprocess_dataframe, merge_continuation_tables
except ImportError:
    from table_postprocess import postprocess_dataframe, merge_continuation_tables


def make_safe_filename(name):
    """Clean string to make it safe for folder/file names."""
    # Replace illegal characters
    safe = re.sub(r'[\\/*?:"<>|]', '_', name)
    # Remove extra spaces
    safe = re.sub(r'\s+', ' ', safe)
    return safe[:120].strip().strip('.') # Limit length to prevent path length issues

def autofit_excel_columns(file_path, title=None):
    """Auto-fit column widths in the output Excel sheet and ensure full macOS QuickLook / Office metadata compatibility."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(file_path)
        
        # Ensure Dublin Core & Office core properties are populated for macOS QuickLook / Finder preview
        wb.properties.creator = "Zotero Table Extractor"
        wb.properties.lastModifiedBy = "Zotero Table Extractor"
        if title:
            wb.properties.title = str(title)[:100]
            
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = 0
                col_letter = col[0].column_letter
                # 采样前 200 行计算列宽，避免超大表格遍历耗时
                for row_idx, cell in enumerate(col):
                    if row_idx >= 200:
                        break
                    if cell.value is not None:
                        lines = str(cell.value).split('\n')
                        for line in lines:
                            max_len = max(max_len, len(str(line)))
                sheet.column_dimensions[col_letter].width = max(min(max_len + 3, 50), 10)
        wb.save(file_path)
    except Exception as e:
        print(f"Warning: Failed to auto-fit column widths: {e}")


def _load_excel_with_attrs(tbl_path, output_path=None):
    """从磁盘加载 Excel 表格，并尽可能恢复其 .attrs 元数据 (标题、页码等)。"""
    existing_df = pd.read_excel(tbl_path)
    if not hasattr(existing_df, 'attrs'):
        existing_df.attrs = {}
    try:
        from openpyxl import load_workbook
        wb = load_workbook(tbl_path, read_only=True)
        if wb.properties and wb.properties.title:
            existing_df.attrs['table_title'] = wb.properties.title
            existing_df.attrs['label'] = wb.properties.title
        wb.close()
    except Exception:
        pass
    if output_path:
        cap_file = os.path.join(output_path, "table_captions.txt")
        if os.path.exists(cap_file):
            base = os.path.splitext(os.path.basename(tbl_path))[0]
            try:
                with open(cap_file, 'r', encoding='utf-8') as f:
                    for line in f:
                        if ':' in line:
                            k, v = line.split(':', 1)
                            if k.strip() == base or make_safe_filename(k.strip()) == base:
                                existing_df.attrs['table_title'] = v.strip()
                                existing_df.attrs['label'] = k.strip()
                                break
            except Exception:
                pass
    return existing_df


def save_tables_to_excel(dataframes, output_path, headers=None, single_file=True, skip_labels=None, skip_supplementary=False):
    """
    Saves a list of DataFrames to Excel.
    If single_file is True, saves to a single file with multiple sheets (if more than 1 df).
    Otherwise, saves to output_path directory as Table 1.xlsx, Table 2.xlsx, etc.
    """
    if not dataframes:
        return False

    if single_file and not str(output_path).lower().endswith('.xlsx'):
        single_file = False

    # 0. 规范化输入：无缝兼容 ExtractedTable 统一对象、dict 管道与原生 pd.DataFrame
    norm_dfs = []
    for d in dataframes:
        if d is None:
            continue
        if hasattr(d, 'to_dataframe'):
            norm_dfs.append(d.to_dataframe())
        elif hasattr(d, 'df') and isinstance(d.df, pd.DataFrame):
            if hasattr(d, 'sync_attrs'):
                d.sync_attrs()
            norm_dfs.append(d.df)
        elif isinstance(d, dict) and 'df' in d:
            df_item = d['df']
            if hasattr(df_item, 'attrs'):
                for k in ('label', 'table_title', 'caption', 'page_idx', 'extractor', 'source'):
                    if k in d and k not in df_item.attrs:
                        df_item.attrs[k] = d[k]
            norm_dfs.append(df_item)
        elif isinstance(d, pd.DataFrame):
            norm_dfs.append(d)

    if not norm_dfs:
        return False

    # 1. 预先合并跨页未命名续表/同名续表，避免未命名 1 行续表被 postprocess 当作空表过滤
    merged_raw_dfs = merge_continuation_tables(norm_dfs)

    cleaned_dfs = []
    for df in merged_raw_dfs:
        if df is None or df.empty:
            continue
        label = df.attrs.get('label')
        title = df.attrs.get('table_title')
        if skip_supplementary and is_supplementary(label or title):
            print(f"[Filter] 跳过附表: {label or title or '未命名'}")
            continue
        if skip_labels and (label or title):
            check_label = format_table_label(label or title)
            safe_check = make_safe_filename(check_label)
            if safe_check in skip_labels:
                print(f"[Filter] 表格 {safe_check} 已由先前步骤提取，跳过以避免重复")
                continue
        proc = postprocess_dataframe(df, headers)
        if proc is not None and not proc.empty:
            if is_metadata_table(proc):
                print(f"[Filter] 过滤非数据表/元数据表: {proc.attrs.get('table_title') or proc.attrs.get('label') or '未命名'}")
                continue
            is_vlm = (proc.attrs.get('extractor') == 'paddleocr_vl' or proc.attrs.get('source') == 'paddleocr_vl')
            if not is_vlm:
                proc = proc.drop_duplicates(keep='first').reset_index(drop=True)
            cleaned_dfs.append(proc)

    if not cleaned_dfs:
        return False
        
    if single_file:
        # Ensure parent directory exists
        parent_dir = os.path.dirname(output_path)
        if parent_dir and not os.path.exists(parent_dir):
            os.makedirs(parent_dir, exist_ok=True)
            
        if len(cleaned_dfs) == 1:
            cleaned_dfs[0].to_excel(output_path, index=False)
            autofit_excel_columns(output_path)
        else:
            with pd.ExcelWriter(output_path) as writer:
                main_count = 0
                supp_count = 0
                for idx, df in enumerate(cleaned_dfs):
                    label = df.attrs.get('label')
                    title = df.attrs.get('table_title')
                    is_supp = is_supplementary(label or title)
                    
                    if label:
                        sheet_name = re.sub(r'[\\/*?:\[\]]', '_', label)[:31].strip()
                    else:
                        if is_supp:
                            supp_count += 1
                            sheet_name = f"Supplementary Table {supp_count}"
                        else:
                            main_count += 1
                            sheet_name = f"Table {main_count}"
                            
                    # De-duplicate sheet names
                    orig_sheet_name = sheet_name
                    suffix = 1
                    while sheet_name in writer.sheets:
                        sheet_name = f"{orig_sheet_name[:27]}_{suffix}"
                        suffix += 1
                        
                    df.to_excel(writer, sheet_name=sheet_name, index=False)
            autofit_excel_columns(output_path)
        print(f"Saved {len(cleaned_dfs)} tables to Excel file: {output_path}")
    else:
        # Output is a directory, save individual Table X.xlsx files
        if not os.path.exists(output_path):
            os.makedirs(output_path, exist_ok=True)
            
        captions = []
        main_count = 0
        supp_count = 0
        written_files = {}  # tbl_path -> DataFrame (with preserved .attrs)
        for idx, df in enumerate(cleaned_dfs):
            # Try to resolve title/caption and label from attrs
            title = df.attrs.get('table_title') or ""
            label = df.attrs.get('label') or ""
            
            # 语种感知
            is_cn = bool(
                re.search(r'[\u4e00-\u9fa5]', title) or 
                re.search(r'[\u4e00-\u9fa5]', label) or 
                (df.size > 0 and re.search(r'[\u4e00-\u9fa5]', ''.join(str(c) for c in df.columns[:5])))
            )
            
            if not label and not title:
                label = f"表{idx+1}" if is_cn else f"Table {idx+1}"

            has_chapter_labels = any(
                bool(re.match(r'^(?:附表|附录表|表|Supplementary\s+Table|Table)\s*\d+[\-–—\.]\d+', str(d.attrs.get('label', '')).strip(), re.IGNORECASE))
                for d in cleaned_dfs
            )

            # 从 title 或 label 中精准提取文中真实表号
            extracted_label = ""
            # 1. 优先从 title 提取（title 通常包含完整表题如 "表4.1 全区岩石..." 或 "表4-1"）
            if title:
                title_match = TABLE_LABEL_RE.search(title)
                if title_match:
                    extracted_label = format_table_label(title_match.group(0))
            
            # 2. 从 label 提取
            if not extracted_label and label:
                # 检查是否为无章节信息的泛化单数字占位符（如 "Table 1", "表1"）且文档包含多级章节表
                is_generic_placeholder = bool(re.match(r'^(?:Table\s*\d+|表\s*\d+)$', label.strip(), re.IGNORECASE))
                if is_generic_placeholder and has_chapter_labels:
                    extracted_label = ""
                else:
                    label_match = TABLE_LABEL_RE.search(label)
                    if label_match:
                        extracted_label = format_table_label(label_match.group(0))
                    else:
                        extracted_label = format_table_label(label)
                        
            # 3. 过滤 OCR 乱码/噪声表号（如 Table bo 等）
            if extracted_label and re.match(r'^Table\s+[a-z]{2,}$', extracted_label, re.IGNORECASE):
                extracted_label = ""
                
            # 若全文包含多级章节表号（如 表1-1 或 表1.1），过滤掉伪单数字表号（如 表1、表2、Table 1）
            is_supp = is_supplementary(label or title)
            if extracted_label:
                file_label = extracted_label
            else:
                # 若全文普遍使用多级章节表号（如 表1-1、表1.1），无表题的零散表格碎片不单独导出为泛化数字 表1、表2
                if has_chapter_labels:
                    print(f"Skipping unlabelled table fragment on page {df.attrs.get('page_idx')} in chapter document...")
                    continue
                # 语种感知回退（中文文献用 表1、附表1；英文文献用 Table 1、Table S1）
                if is_supp:
                    supp_count += 1
                    file_label = f"附表{supp_count}" if is_cn else f"Table S{supp_count}"
                else:
                    main_count += 1
                    file_label = f"表{main_count}" if is_cn else f"Table {main_count}"
                
            file_label = file_label.replace('\xa0', ' ').strip()
            safe_label = make_safe_filename(file_label)
            
            if skip_labels and safe_label in skip_labels:
                print(f"Table {safe_label} was already successfully extracted online. Skipping PDF version to preserve quality.")
                continue
            
            tbl_path = os.path.join(output_path, f"{safe_label}.xlsx")
            
            # Check if this file already exists (e.g. from an earlier page or part of a continued table)
            if os.path.exists(tbl_path):
                print(f"File {tbl_path} already exists. Merging dataframes...")
                try:
                    if tbl_path in written_files:
                        existing_df = written_files[tbl_path]
                    else:
                        existing_df = _load_excel_with_attrs(tbl_path, output_path)
                    
                    # Determine relationship between new df and existing_df
                    cols1 = [str(c).lower().strip() for c in existing_df.columns]
                    cols2 = [str(c).lower().strip() for c in df.columns]
                    
                    compare_cols1 = [c for c in cols1 if 'unnamed' not in c]
                    compare_cols2 = [c for c in cols2 if 'unnamed' not in c]
                    
                    intersection = set(compare_cols1).intersection(set(compare_cols2))
                    overlap_ratio = len(intersection) / max(len(compare_cols1), len(compare_cols2)) if compare_cols1 and compare_cols2 else 0

                    # 判断是否为同一张大表的跨页续表部分（续表应纵向拼接）
                    title_str = str(df.attrs.get('table_title', '')).lower()
                    is_cont_part = (
                        bool(re.search(r'(?:续表|接上表|continued|\(cont\)|cont\.)', title_str)) or
                        '(续)' in str(df.attrs.get('table_title', '')) or
                        bool(df.attrs.get('is_continuation'))
                    )
                    
                    if is_cont_part:
                        from table_postprocess import combine_df_group
                        print("-> Detected continuation fragment. Merging with existing table...")
                        df_to_save = combine_df_group([existing_df, df])
                    elif overlap_ratio >= 0.5:
                        existing_valid_cells = existing_df.dropna(how='all').notna().sum().sum()
                        new_valid_cells = df.dropna(how='all').notna().sum().sum()
                        
                        # 若已有表经过跨页多页拼合且信息量显著更多，保护多页合并版本
                        if existing_valid_cells > new_valid_cells * 1.4 and len(existing_df) > len(df) * 1.3:
                            print("-> Keeping larger multi-page existing version...")
                            df_to_save = existing_df
                        elif df.attrs.get('extractor') in ('paddleocr_vl', 'ppstructure_vlm'):
                            print("-> High-priority OCR table replaces existing version...")
                            df_to_save = df
                        else:
                            df_to_save = df if new_valid_cells >= existing_valid_cells else existing_df
                    else:
                        # 检查是否为同名但位于不同章节/页面的独立数据表（如论文作者笔误重复表号）
                        curr_p = df.attrs.get('page_idx')
                        existing_p = existing_df.attrs.get('page_idx')
                        is_distant_table = False
                        if curr_p is not None and existing_p is not None and abs(curr_p - existing_p) >= 2:
                            is_distant_table = True
                        elif curr_p is not None and curr_p > 5 and existing_p is None:
                            # 现有表缺少页码属性但当前表在后方页面，保留为独立表
                            is_distant_table = True

                        if is_distant_table and curr_p is not None:
                            # 独立表格另存为带页码的规范文件名，避免覆盖前表
                            tbl_path = os.path.join(output_path, f"{safe_label}_P{curr_p + 1}.xlsx")
                            df_to_save = df
                            print(f"-> Detected distinct table on page {curr_p + 1} sharing label {safe_label}. Saving as {tbl_path}...")
                        elif df.attrs.get('extractor') in ('paddleocr_vl', 'ppstructure_vlm') or df.size >= existing_df.size:
                            print("-> High-priority OCR table replaces lower-priority version...")
                            df_to_save = df
                        else:
                            print("-> Keeping existing higher-quality table...")
                            df_to_save = existing_df
                            
                    # Drop duplicate rows if any got appended/merged while strictly preserving natural top-to-bottom row order
                    try:
                        df_to_save = df_to_save.drop_duplicates(keep='first').reset_index(drop=True)
                    except Exception as e:
                        print(f"Warning during duplicate row cleaning: {e}")
                        
                    # 严格保留并合并 .attrs 元数据，避免覆写丢失
                    saved_attrs = dict(getattr(existing_df, 'attrs', {}))
                    saved_attrs.update({k: v for k, v in getattr(df, 'attrs', {}).items() if v is not None})
                    if not hasattr(df_to_save, 'attrs') or df_to_save.attrs is None:
                        df_to_save.attrs = {}
                    df_to_save.attrs.update(saved_attrs)
                        
                    df_to_save.to_excel(tbl_path, index=False)
                    autofit_excel_columns(tbl_path, title=df_to_save.attrs.get('table_title') or title or file_label)
                    written_files[tbl_path] = df_to_save
                    print(f"Saved: {tbl_path}")
                except Exception as e:
                    print(f"Error merging dataframes for {tbl_path}: {e}")
            else:
                df.to_excel(tbl_path, index=False)
                autofit_excel_columns(tbl_path, title=title or file_label)
                written_files[tbl_path] = df
                print(f"Saved: {tbl_path}")
            
            # Collect caption info
            full_title = title.strip() if title and title.strip() else file_label
            captions.append((file_label, full_title))
            
        # Write table_captions.txt in the same directory
        captions_path = os.path.join(output_path, "table_captions.txt")
        best_captions = {}
        if os.path.exists(captions_path):
            try:
                with open(captions_path, "r", encoding="utf-8") as f_ex:
                    for line in f_ex:
                        line_s = line.strip()
                        if line_s and not line_s.startswith("#") and "\t" in line_s:
                            k, v = line_s.split("\t", 1)
                            if v.strip():
                                best_captions[k.strip()] = v.strip()
            except Exception:
                pass

        # Also inspect existing .xlsx files on disk to ensure all present tables are listed
        import glob
        for fpath in glob.glob(os.path.join(output_path, "*.xlsx")):
            fname = os.path.splitext(os.path.basename(fpath))[0]
            if fname not in best_captions:
                best_captions[fname] = fname

        for lbl, cap in captions:
            if lbl not in best_captions or (len(cap) > len(best_captions[lbl]) and cap != lbl):
                best_captions[lbl] = cap
                
        # 仅保留磁盘上实际存在的 .xlsx 文件对应表头
        best_captions = {k: (v if v.strip() else k) for k, v in best_captions.items() if os.path.exists(os.path.join(output_path, f"{k}.xlsx"))}

        if best_captions:
            # Convert back to natural sorted list
            sorted_captions = list(best_captions.items())
            try:
                def natural_sort_key(item):
                    return [int(c) if c.isdigit() else c.lower() for c in re.split(r'(\d+)', str(item[0]))]
                sorted_captions.sort(key=natural_sort_key)
            except Exception:
                pass
                
            # 自动检测表号不连续与跳号异常（例如 表6-4 到 表6-6 之间跳过 表6-5）
            anomalies = []
            chapter_groups = {}
            for lbl, cap in sorted_captions:
                m_chap = re.search(r'(\d+)[\-–—\.](\d+)', lbl)
                m_single = re.search(r'^(?:表|Table\s*)(\d+)$', lbl, re.IGNORECASE)
                if m_chap:
                    chap, sec = int(m_chap.group(1)), int(m_chap.group(2))
                    chapter_groups.setdefault(chap, []).append((sec, lbl))
                elif m_single:
                    chapter_groups.setdefault(0, []).append((int(m_single.group(1)), lbl))

            for chap, items in chapter_groups.items():
                items.sort(key=lambda x: x[0])
                nums = [x[0] for x in items]
                for i in range(len(nums) - 1):
                    if nums[i+1] - nums[i] > 1:
                        missing = [f"第{chap}章 表{chap}-{x}" if chap > 0 else f"表{x}" for x in range(nums[i]+1, nums[i+1])]
                        anomalies.append(f"【表序跳号注意】: 存在不连续表号，缺失: {', '.join(missing)}")

            captions_path = os.path.join(output_path, "table_captions.txt")
            try:
                with open(captions_path, "w", encoding="utf-8") as f:
                    if anomalies:
                        f.write("# ========================================================\n")
                        f.write("# 【表序号与排版完整性诊断提示】：\n")
                        for a in anomalies:
                            f.write(f"# {a}\n")
                        f.write("# ========================================================\n\n")
                    for lbl, cap in sorted_captions:
                        # Tab-separated: Label \t Title
                        f.write(f"{lbl}\t{cap}\n")
                print(f"Saved table captions to: {captions_path}")
            except Exception as e:
                print(f"Warning: Failed to save captions file: {e}")
            
    return True
